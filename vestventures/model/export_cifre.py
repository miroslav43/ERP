#!/usr/bin/env python3
"""
Reads RECALCULATED workbooks (cached values written by LibreOffice) and exports numbers.

    python3 model/export_cifre.py snapshot  model/recalc/scen1.xlsx model/recalc/scen2.xlsx model/recalc/scen3.xlsx \
                                            --out model/recalc/snapshot.json
    python3 model/export_cifre.py cifre     financial-model.xlsx model/recalc/snapshot.json \
                                            --out _surse/cifre.json

Every value comes from the workbook's cached results (openpyxl data_only=True), never from a
re-implementation of the model.  A missing cached value aborts the export.
"""

from __future__ import annotations

import argparse
import datetime as dt
import json
import sys
from pathlib import Path

from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parent))
import build_model as bm  # noqa: E402

SCEN = bm.SCEN_NAMES


class Book:
    def __init__(self, path: str):
        self.wb = load_workbook(path, data_only=True)
        self.path = path

    def cell(self, sheet: str, coord: str):
        v = self.wb[sheet][coord].value
        if v is None:
            raise SystemExit(f"{self.path}: {sheet}!{coord} has no cached value - was the workbook recalculated?")
        return v

    def name(self, nm: str):
        dn = self.wb.defined_names[nm]
        (sheet, coord), = list(dn.destinations)
        return self.cell(sheet, coord.replace("$", ""))


def num(v, nd=0):
    if isinstance(v, (int, float)):
        return round(v, nd) if nd else int(round(v))
    return v  # text such as "> 36" or "not within 36 months"


def month_label(book: Book, m):
    if not isinstance(m, (int, float)):
        return None
    d = book.cell("Revenue", f"{bm.col(int(m))}4")
    if isinstance(d, (dt.datetime, dt.date)):
        return d.strftime("%b %Y")
    return None


def scenario_metrics(b: Book) -> dict:
    R = lambda m, row: b.cell("Revenue", f"{bm.col(m)}{row}")
    out = {
        "scenario": b.name("ScenarioIdx"),
        "customers_m12": num(R(12, 21)), "mrr_m12_eur": num(R(12, 26)), "arr_m12_eur": num(R(12, 27)),
        "customers_m24": num(R(24, 21)), "mrr_m24_eur": num(R(24, 26)), "arr_m24_eur": num(R(24, 27)),
        "customers_m36": num(R(36, 21)), "mrr_m36_eur": num(R(36, 26)), "arr_m36_eur": num(R(36, 27)),
        "breakeven_month": num(b.name("in_Breakeven")),
        "cash_low_point_eur": num(b.name("in_CashLow")),
        "cash_low_point_month": num(b.name("in_CashLowMonth")),
        "cash_m36_eur": num(b.name("in_Cash36")),
        "net_cash_flow_m36_eur": num(b.cell("P&L & Cash", f"{bm.col(36)}18")),
        "customers_m36_exact": round(R(36, 21), 1),
        "runway_zero_revenue_months": num(b.name("in_RunwayZeroRev")),
        "runway_plan_months": num(b.name("in_RunwayPlan")),
        "months_to_30": num(b.name("in_M30")),
        "months_to_100": num(b.name("in_M100")),
        "months_to_300": num(b.name("in_M300")),
        "revenue_y1_eur": num(b.name("in_Rev12")), "revenue_y2_eur": num(b.name("in_Rev24")), "revenue_y3_eur": num(b.name("in_Rev36")),
        "costs_y1_eur": num(b.name("in_Cost12")), "costs_y2_eur": num(b.name("in_Cost24")), "costs_y3_eur": num(b.name("in_Cost36")),
    }
    out["breakeven_date"] = month_label(b, out["breakeven_month"])
    for n in (30, 100, 300):
        out[f"date_to_{n}"] = month_label(b, out[f"months_to_{n}"])
    return out


def round_tests(b: Book) -> list:
    ws = b.wb["Use of funds"]
    out = []
    for r in range(1, ws.max_row + 1):
        if ws[f"A{r}"].value == "Round size tested":
            rr = r + 1
            while ws[f"A{rr}"].value and ws[f"B{rr}"].value is not None:
                out.append({"case": ws[f"A{rr}"].value, "round_eur": num(b.cell("Use of funds", f"B{rr}")),
                            "runway_zero_revenue_months": num(b.cell("Use of funds", f"C{rr}")),
                            "runway_plan_months": num(b.cell("Use of funds", f"D{rr}"))})
                rr += 1
    return out


def cmd_snapshot(a):
    snap = {}
    for path, name in zip(a.books, SCEN):
        b = Book(path)
        m = scenario_metrics(b)
        if m["scenario"] != SCEN.index(name) + 1:
            raise SystemExit(f"{path}: scenario selector is {m['scenario']}, expected {name}")
        snap[name] = m
    Path(a.out).write_text(json.dumps(snap, indent=2, ensure_ascii=False))
    print(json.dumps(snap, indent=2, ensure_ascii=False))


def cmd_cifre(a):
    fm = Book(a.books[0])
    snap = json.loads(Path(a.books[1]).read_text())
    base = scenario_metrics(fm)
    # the delivered workbook must reproduce the base snapshot exactly
    for k, v in snap["Base"].items():
        if base.get(k) != v:
            raise SystemExit(f"base mismatch on {k}: workbook {base.get(k)} vs snapshot {v}")
    cover = fm.wb["Cover"]
    integ = [r for r in range(1, cover.max_row + 1) if str(cover[f"B{r}"].value or "").startswith("Integrity checks")]
    if len(integ) != 1 or fm.cell("Cover", f"C{integ[0]}") != "OK":
        raise SystemExit("Cover integrity check is not OK")

    fx = fm.name("in_FX")
    round_eur = num(fm.name("in_Round"))
    uof = []
    ws = fm.wb["Use of funds"]
    first, last = 9, 15
    for r in range(first, last + 1):
        uof.append({"category": ws[f"A{r}"].value, "eur": num(fm.cell("Use of funds", f"C{r}")),
                    "pct": round(fm.cell("Use of funds", f"D{r}") * 100, 1)})
    # rounding fix: make euros sum exactly to the round
    diff = round_eur - sum(x["eur"] for x in uof)
    big = max(range(len(uof)), key=lambda i: uof[i]["eur"])
    uof[big]["eur"] += diff
    uof[big]["pct"] = round(uof[big]["pct"] + 100 - sum(x["pct"] for x in uof), 1)

    pricing = []
    for key, label, ron, mix in bm.PRICES_RON:
        pricing.append({"plan": label, "ron_month": ron, "eur_month": round(ron / fx, 1), "plan_mix_pct": round(mix * 100),
                        "source": "src/content/landing/preturi.ts"})
    pricing.append({"note": "Flat price per company up to 20 employees; first month free; no setup fee; seller not "
                            "VAT-registered so prices are final. No published price above 20 employees. Add-on modules "
                            "15-69 RON/month each (payroll 69). Accountant pilot: free until 31 Mar 2027, then 20% "
                            "commission for 6 months to the accountant."})

    snm = {k: num(fm.name(f"in_{k}")) for k in ("SalesStart", "SalesFullStart", "SupportStart", "SupportFullStart", "DevStart", "FounderPayStart",
                                                 "PilotConvMonth", "PayrollValEnd", "A11yEnd", "PentestEnd")}
    fte = {k: round(fm.name(f"in_{k}"), 2) for k in ("SalesFTE", "SupportFTE")}
    lab = lambda m: f"M{m} ({month_label(fm, m)})" if isinstance(m, int) else str(m)
    milestones = [
        {"month": lab(1), "milestone": "Pre-seed round closes (assumed); both co-founders full time, unpaid"},
        {"month": lab(snm["FounderPayStart"]), "milestone": "Co-founders start drawing a salary"},
        {"month": lab(snm["SalesStart"]), "milestone": f"Partner & sales manager hired (accountant channel), part-time ({fte['SalesFTE']:.0%})"},
        {"month": lab(snm["PilotConvMonth"]), "milestone": "Accountant pilot ends (free until 31 Mar 2027); pilot companies convert to paid (founders onboard them)"},
        {"month": lab(snm["SupportStart"]), "milestone": f"Customer onboarding & support specialist hired, part-time ({fte['SupportFTE']:.0%})"},
        {"month": lab(snm["SalesFullStart"]), "milestone": "Partner & sales manager full time"},
        {"month": lab(snm["PayrollValEnd"]), "milestone": "Payroll legal values validated by an accountant (payroll moves from 'built, undergoing validation' to sellable)"},
        {"month": lab(snm["A11yEnd"]), "milestone": "Accessibility audit (EN 301 549) completed"},
        {"month": lab(snm["PentestEnd"]), "milestone": "External penetration test (after the seed)"},
        {"month": lab(base["months_to_30"]), "milestone": "30 paying customers"},
        {"month": lab(snm["DevStart"]), "milestone": "Developer hired (after the seed)"},
        {"month": lab(snm["SupportFullStart"]), "milestone": "Onboarding & support specialist full time (after the seed)"},
        {"month": lab(base["months_to_100"]), "milestone": "100 paying customers - seed-readiness trigger (start seed / VV Seed conversations)"},
    ]
    if isinstance(base["breakeven_month"], int):
        milestones.append({"month": lab(base["breakeven_month"]), "milestone": "Monthly break-even (net cash flow >= 0)"})
    if isinstance(base["months_to_300"], int):
        milestones.append({"month": lab(base["months_to_300"]), "milestone": "300 paying customers"})
    milestones.sort(key=lambda x: int(x["month"].split()[0][1:]))

    arpa = fm.name("in_ARPA")
    som_arr = base["arr_m36_eur"]
    cifre = {
        "_meta": {
            "generated": "2026-10-07",
            "generated_from": ["financial-model.xlsx (recalculated by LibreOffice)", "model/recalc/snapshot.json"],
            "rule": "Single source of numbers for every VestVentures document. Do not type numbers by hand elsewhere.",
            "currency": "EUR; RON converted at 5.0 RON = 1 EUR",
            "status": "Everything after 6 Oct 2026 is a PROJECTION. Zero paying customers today.",
        },
        "round_eur": round_eur,
        "round_investor": {"name": "Vest Ventures", "programme": "Accelerator",
                           "vest_ventures_ticket_eur": num(fm.name("in_VVTicket")),
                           "note": "Vest Ventures is the only investor in the pre-seed round; ticket inside VV's published EUR 10k-200k accelerator range"},
        "stage": "Pre-seed (apply to the Vest Ventures Accelerator, not Seed: zero revenue)",
        "instrument": "Convertible loan agreement (CLA) or SHA on Vest Ventures' templates; a CLA would use the pre-money valuation as its cap",
        "pre_money_eur": num(fm.name("in_PreMoney")),
        "post_money_eur": num(fm.name("in_PostMoney")),
        "runway_months": base["runway_zero_revenue_months"],
        "runway_basis": "Months the round pays every planned cost with ZERO revenue (base cost plan)",
        "runway_plan_months": base["runway_plan_months"],
        "round_size_test": round_tests(fm),
        "use_of_funds": uof,
        "use_of_funds_basis": f"Share of planned spend over the zero-revenue runway (months 1-{base['runway_zero_revenue_months']}), applied to the round",
        "close_month": "Jan 2027 (model month 1) - ASSUMPTION; Vest Ventures Accelerator Cohort 3 dates are not published",
        "model_start": "Jan 2027",
        "base": base,
        "base_note": ("Base case is NOT default-alive: without a seed round cash runs out after runway_plan_months; "
                      "the seed raise starts at the 100-paying-customer milestone. cash_low_point_eur < 0 = shortfall without a seed."),
        "conservative": snap["Conservative"],
        "upside": snap["Upside"],
        "unit_economics": {
            "basis": "Base scenario, model year 2 (Jan-Dec 2028)",
            "arpa_eur": round(arpa, 2),
            "gross_margin_pct": round(fm.name("in_CM") * 100, 1),
            "gross_margin_definition": "Contribution margin per customer: (ARPA - variable infra - payment fees) / ARPA",
            "gross_margin_pnl_pct": round(fm.name("in_GMPL") * 100, 1),
            "gross_margin_pnl_definition": "P&L gross margin after fixed hosting and the onboarding/support salary (year 2)",
            "cac_eur": num(fm.name("in_CAC")),
            "cac_definition": "Fully loaded: partner manager salary + marketing + travel + accountant commissions / new paying customers",
            "cac_paid_ads_eur": num(fm.name("in_CACPaid")),
            "monthly_churn_pct": round(fm.name("in_ChurnUE") * 100, 1),
            "ltv_eur": num(fm.name("in_LTV")),
            "ltv_cac": round(fm.name("in_LTVCAC"), 1),
            "payback_months": round(fm.name("in_Payback"), 1),
            "benchmarks": "ChartMogul median churn 6.1%/month for ARPA < US$25; SMB CAC payback median ~11 months, 6-14 healthy (piata.md §d.1)",
        },
        "team_plan": [
            {"role": "Co-founder & CEO - Miroslav Maletici", "from_month": 1, "paid_from_month": snm["FounderPayStart"],
             "gross_eur_month": num(fm.name("in_SalFounder")), "note": "full time from close, unpaid until paid_from_month"},
            {"role": "Co-founder & CTO - Răzvan Pervulescu", "from_month": 1, "paid_from_month": snm["FounderPayStart"],
             "gross_eur_month": num(fm.name("in_SalFounder")), "note": "full time from close, unpaid until paid_from_month"},
            {"role": "Partner & sales manager (accountant channel)", "from_month": snm["SalesStart"], "gross_eur_month": num(fm.name("in_SalSales")),
             "gross_basis": "full-time equivalent", "fte_until_full_time": fte["SalesFTE"], "full_time_from_month": snm["SalesFullStart"],
             "note": "new accountant partners per month scale with this FTE"},
            {"role": "Customer onboarding & support specialist", "from_month": snm["SupportStart"], "gross_eur_month": num(fm.name("in_SalSupport")),
             "gross_basis": "full-time equivalent", "fte_until_full_time": fte["SupportFTE"], "full_time_from_month": snm["SupportFullStart"]},
            {"role": "Full-stack developer", "from_month": snm["DevStart"], "gross_eur_month": num(fm.name("in_SalDev")),
             "note": "hired after the seed; not funded by the pre-seed round"},
        ],
        "pricing": pricing,
        "milestones": milestones,
        "tam_sam_som": {
            "fx": "5 RON = 1 EUR",
            "tam": {"eur_year": 188_000_000, "ron_year": 940_488_000,
                    "formula": "526,000 employers with >=1 employee x 149 RON x 12",
                    "label": "DERIVED",
                    "sources": ["ICAP CRIF FY2025 via https://curierulnational.ro/peste-jumatate-de-milion-de-firme-aveau-388-milioane-salariati-in-2025/",
                                "price: src/content/landing/preturi.ts"]},
            "sam": {"eur_year": 58_000_000, "ron_year": 289_926_696,
                    "formula": "29,461 firms (10-19 persons) x 398 RON x 12 + 24,920 firms (20-249) x 499 RON x 12",
                    "label": "DERIVED; understated because no price above 20 employees is published",
                    "sources": ["Eurostat SBS 2023 sbs_sc_ovw https://ec.europa.eu/eurostat/databrowser/view/sbs_sc_ovw/default/table"]},
            "sam_extension_accountants": {"eur_year": 16_800_000, "ron_year": 84_027_060,
                                          "formula": "~470,000 micro-employers (DERIVED estimate) x 10% reachable via accountants (ASSUMPTION) x 149 RON x 12",
                                          "label": "DERIVED + ASSUMPTION; show separately, never add silently"},
            "som": {"eur_arr": som_arr,
                    "formula": f"base-case paying customers at month 36 ({base['customers_m36']}) x ARPA x 12 = MRR month 36 x 12",
                    "label": "PROJECTION (financial model, base scenario, Dec 2029)",
                    "share_of_employers_pct": round(base["customers_m36"] / 526_000 * 100, 3),
                    "conservative_eur_arr": snap["Conservative"]["arr_m36_eur"], "upside_eur_arr": snap["Upside"]["arr_m36_eur"],
                    "note": "piata.md §e illustrated 600 firms x 250 RON x 12 = EUR 360k ARR; the model's bottom-up base case replaces it"},
        },
        "benchmarks_for_ask": {
            "romania_avg_pre_seed_2025_eur": 526_000, "romania_avg_seed_2025_eur": 1_500_000,
            "source": "How to Web & Underline Ventures, Venture in Eastern Europe 2025 (piata.md §d.2)",
            "vv_accelerator_ticket": "EUR 10k-200k (vestventures.md §4)",
            "valuation_benchmark": "No sourced RO/CEE pre-seed valuation benchmark (UNVERIFIED)",
        },
    }
    Path(a.out).write_text(json.dumps(cifre, indent=2, ensure_ascii=False) + "\n")
    print(json.dumps(cifre, indent=2, ensure_ascii=False))


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("mode", choices=["snapshot", "cifre"])
    ap.add_argument("books", nargs="+")
    ap.add_argument("--out", required=True)
    a = ap.parse_args()
    cmd_snapshot(a) if a.mode == "snapshot" else cmd_cifre(a)


if __name__ == "__main__":
    main()
