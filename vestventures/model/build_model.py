#!/usr/bin/env python3
"""
Generates the ADMINISTRATIVO investor financial model and the standalone cap table.

    python3 model/build_model.py                 # writes the workbooks (formulas only)
    python3 model/build_model.py --snapshot f.json  # also writes the scenario snapshot table

Outputs (relative to /srv/apps/ERP/vestventures):
    financial-model.xlsx   live Excel formulas; every driver lives on "Assumptions"
                           or "Scenarios"; Cover!C8 selects the scenario (1/2/3)
    cap-table.xlsx         today / pre-seed / illustrative seed + CLA alternative

Recalculation (cached values for readers that do not recalculate) is done by
model/recalc.sh with LibreOffice in a container; model/export_cifre.py then
reads the recalculated workbooks back and writes _surse/cifre.json.

Conventions: blue font = hard-coded input, green = link to another sheet,
black = formula, yellow fill = key lever.  EUR throughout, RON converted at
in_FX (5.0 RON = 1 EUR).  Figures marked PROJECTION are projections, not facts.
"""

from __future__ import annotations

import argparse
import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName

ROOT = Path(__file__).resolve().parent.parent
N_MONTHS = 36
FIRST_COL = 3  # month 1 lives in column C
LAST_COL = FIRST_COL + N_MONTHS - 1  # column AL

FONT = "Arial"
BLUE = Font(name=FONT, color="0000FF")
BLUE_B = Font(name=FONT, color="0000FF", bold=True)
GREEN = Font(name=FONT, color="008000")
BLACK = Font(name=FONT, color="000000")
BOLD = Font(name=FONT, bold=True)
TITLE = Font(name=FONT, bold=True, size=14)
SUB = Font(name=FONT, italic=True, color="555555", size=9)
HDR = Font(name=FONT, bold=True, color="FFFFFF")
HDR_FILL = PatternFill("solid", fgColor="1F3864")
SEC_FILL = PatternFill("solid", fgColor="D9E1F2")
YELLOW = PatternFill("solid", fgColor="FFFF00")
THIN = Side(style="thin", color="999999")
TOP = Border(top=THIN)

EUR = '"€"#,##0;("€"#,##0);"-"'
EUR2 = '"€"#,##0.00;("€"#,##0.00);"-"'
RON = '#,##0" RON";(#,##0" RON");"-"'
NUM = '#,##0;(#,##0);"-"'
NUM1 = '#,##0.0;(#,##0.0);"-"'
PCT = '0.0%;(0.0%);"-"'
MULT = '0.0"x"'
MON = '0" mo"'

SCEN_NAMES = ["Base", "Conservative", "Upside"]


def col(m: int) -> str:
    """Column letter of model month m (1-based)."""
    return get_column_letter(FIRST_COL + m - 1)


C1, CN = col(1), col(N_MONTHS)


def rng(sheet: str, row: int) -> str:
    return f"'{sheet}'!${C1}${row}:${CN}${row}" if " " in sheet else f"{sheet}!${C1}${row}:${CN}${row}"


def ref(sheet: str, cell: str) -> str:
    return f"'{sheet}'!{cell}" if " " in sheet or "&" in sheet else f"{sheet}!{cell}"


def style_all(ws):
    for row in ws.iter_rows():
        for c in row:
            if c.font is None or c.font.name != FONT:
                f = c.font
                c.font = Font(name=FONT, bold=f.bold, italic=f.italic, color=f.color, size=f.size)


# --------------------------------------------------------------------------------------
# Inputs.  Every hard-coded number of the model is in one of these two tables.
# --------------------------------------------------------------------------------------

# (name, label, Base, Conservative, Upside, unit, number format, note)
SCENARIO_DRIVERS = [
    ("PilotPractices", "Accountant practices that join the pilot", 6, 3, 10, "practices", NUM,
     "Pilot offers 10 places, sign-up until 15 Nov 2026 (administrativo.ro/pentru-contabili). Zero signed on 6 Oct 2026 -> base assumes 6. ASSUMPTION."),
    ("PilotFirmsPerPractice", "Client companies per pilot practice", 2, 1.5, 2.5, "firms", NUM1,
     "Pilot allows 1-3 client companies per practice (administrativo.ro/pentru-contabili). ASSUMPTION within that range."),
    ("PilotConv", "Pilot company -> paying customer (Apr 2027)", 0.50, 0.33, 0.70, "%", PCT,
     "Pilot is free until 31 Mar 2027 (administrativo.ro/pentru-contabili). Conversion not observed yet. ASSUMPTION."),
    ("NewPartnersY1", "New accountant partners per month, year 1 (full-time partner manager)", 2, 1, 3, "per month", NUM1,
     "Signed by the partner manager from in_PartnerStart; rate for a FULL-TIME manager, scaled by in_SalesFTE while the role is part-time "
     "(until in_SalesFullStart). ASSUMPTION (~18,400 practices in RO: CECCAR 2025 report, as reported)."),
    ("NewPartnersY2", "New accountant partners per month, year 2", 3, 1.5, 5, "per month", NUM1, "ASSUMPTION."),
    ("NewPartnersY3", "New accountant partners per month, year 3", 4, 2, 6, "per month", NUM1, "ASSUMPTION."),
    ("RefRate", "Firms referred per active partner per month", 0.20, 0.12, 0.30, "firms", '0.00',
     "0.20 = one referral every 5 months per partner. No sourced figure for clients per accountant. ASSUMPTION."),
    ("RefConv", "Referred firm -> paying customer", 0.50, 0.40, 0.60, "%", PCT, "ASSUMPTION (warm referral by the firm's own accountant)."),
    ("OrganicY1", "Organic / SEO / direct trials per month, year 1", 4, 2, 8, "trials", NUM1,
     "Web traffic today ~7 genuine visitors/month, 201 GSC impressions since 2 Sep 2026 (pre-launch). ASSUMPTION."),
    ("OrganicY2", "Organic / SEO / direct trials per month, year 2", 12, 6, 20, "trials", NUM1, "ASSUMPTION."),
    ("OrganicY3", "Organic / SEO / direct trials per month, year 3", 20, 10, 35, "trials", NUM1, "ASSUMPTION."),
    ("TrialConv", "Trial -> paying customer (after free month)", 0.25, 0.15, 0.30, "%", PCT,
     "First month free for every configuration (administrativo.ro price list). ASSUMPTION, no observed data."),
    ("CostPerTrial", "Paid marketing cost per trial", 80, 120, 60, "EUR", EUR, "Google/LinkedIn, REGES-Online keywords. ASSUMPTION."),
    ("Churn", "Monthly logo churn", 0.03, 0.05, 0.02, "% / month", PCT,
     "ChartMogul median 6.1%/mo for ARPA < US$25; our ARPA ~US$50 sits in a better band. 3% base / 5% conservative / 2% upside: ASSUMPTION, to be tested in the pilot."),
    ("PriceInc", "Annual list-price / upsell increase (from year 2)", 0.05, 0.00, 0.08, "% / year", PCT, "ASSUMPTION."),
    ("DevStart", "Month the developer is hired", 25, 31, 19, "model month", NUM,
     "AFTER the seed: seed raise starts at 100 paying customers (in_M100), ~5-6 months to close. Base Jan 2029; Upside Jul 2028 "
     "(100 customers sooner); Conservative Jul 2029 (seed trigger not reached - deliberately late). Not funded by the pre-seed. ASSUMPTION."),
]

# Plain inputs: (name, label, value, unit, number format, note).  Section rows are ("#", title).
PRICES_RON = [  # (name, label, RON/month, default plan mix)
    ("Core", "HR core (attendance, leave, personnel file, employee portal)", 149, 0.45),
    ("ExtHR", "Extended HR package (core + REGES, onboarding, courses, SSM, evaluations, KPI)", 249, 0.25),
    ("Ops", "Operations package (core + fleet, maintenance, inventory, announcements, ticketing)", 229, 0.10),
    ("Fin", "Finance package (core + payroll, per diem, reports)", 219, 0.05),
    ("All", "Whole application (all 15 add-ons)", 499, 0.15),
]

INPUTS = [
    ("#", "General"),
    ("FX", "Exchange rate RON per EUR", 5.0, "RON/EUR", '0.00', "Convention used in every document of this application."),
    ("StartDate", "Model month 1", "2027-01-01", "date", "mmm yyyy",
     "Model starts Jan 2027 = assumed close of the round. VV Accelerator Cohort 3 date is NOT published (vestventures.vc, 6 Oct 2026) -> if the close slips, the whole timeline shifts."),
    ("FundingMonth", "Month the round is received", 1, "model month", NUM, "ASSUMPTION; see StartDate."),
    ("OpeningCash", "Cash in the company before the round", 0, "EUR", EUR, "[TO CONFIRM] by the founders. Set to 0 (prudent)."),
    ("#", "Round (key levers)"),
    ("VVTicket", "Vest Ventures accelerator ticket", 135000, "EUR", EUR,
     "VV Accelerator ticket EUR 10k-200k, de minimis, >=10% independent private co-investment (vestventures.vc/en/programs-terms). "
     "EUR 135k requested, inside the range."),
    ("Angels", "Independent private co-investors (business angels)", 15000, "EUR", EUR,
     "10% of the round (11.1% of the VV ticket): meets VV's >=10% private co-investment rule whether it is measured on the total "
     "(exactly, no margin) or on VV's ticket - basis [TO CONFIRM] with VV. Angels NOT identified yet."),
    ("MinPrivatePct", "Minimum private co-investment (VV rule, as % of VV ticket)", 0.10, "%", PCT,
     "Vest Ventures programme terms ('>=10% private co-investment'). Basis of the 10% (ticket vs total) [TO CONFIRM] with VV."),
    ("PreMoney", "Pre-money valuation (fully diluted, incl. new ESOP)", 1250000, "EUR", EUR,
     "Founders' proposal. We found no sourced RO/CEE pre-seed valuation benchmark; DERIVED range EUR 0.75-3M post."),
    ("ESOP", "Employee option pool, % of post-money (created pre-money)", 0.10, "%", PCT,
     "Created before the round so it dilutes founders only (standard 'pre-money pool'). SRL needs a phantom/virtual scheme until SRL->SA [TO CONFIRM]."),
    ("#", "Pricing (list prices, RON/month, final - seller not VAT-registered)"),
] + [
    (f"Price{k}", lbl, p, "RON/month", RON, "Public price list, administrativo.ro/en/preturi") for k, lbl, p, _ in PRICES_RON
] + [
    ("#", "Plan mix of paying customers (must sum to 100%)"),
] + [
    (f"Mix{k}", f"Share on {lbl.split(' (')[0]}", m, "%", PCT,
     "ASSUMPTION. Payroll (Finance) kept low: payroll is built, undergoing accountant validation." if k == "Fin" else "ASSUMPTION.")
    for k, lbl, _, m in PRICES_RON
] + [
    ("#", "Channels and timing"),
    ("PilotConvMonth", "Month pilot companies start paying", 4, "model month", NUM, "Apr 2027: pilot is free until 31 Mar 2027 (administrativo.ro/pentru-contabili)."),
    ("PartnerStart", "First month new partners are signed", 4, "model month", NUM, "One month after the partner manager starts (in_SalesStart)."),
    ("PaidShare", "Share of marketing budget spent on performance ads", 0.6, "%", PCT, "Rest is content/SEO/events material. ASSUMPTION."),
    ("CommRate", "Accountant commission on referred subscription", 0.20, "%", PCT, "20% for 6 months from first paid month (administrativo.ro/pentru-contabili)."),
    ("CommMonths", "Commission duration", 6, "months", NUM, "administrativo.ro/pentru-contabili."),
    ("#", "Team (gross monthly salary, EUR; employer adds CAM)"),
    ("CAM", "Employer contribution on gross (CAM)", 0.0225, "%", PCT,
     "Romanian 'contribuția asiguratorie pentru muncă' 2.25% - legal value [TO CONFIRM] by accountant."),
    ("Raise", "Annual salary increase (from year 2)", 0.05, "%", PCT, "ASSUMPTION."),
    ("SalFounder", "Co-founder gross salary (each, 2 people)", 1000, "EUR/month", EUR,
     "Both co-founders full time from close, on a deliberately lean pre-seed salary (~RON 5,000 gross) to stretch the runway; "
     "kept at this level in the model (only the annual raise applies). ASSUMPTION - no sourced salary benchmark."),
    ("SalSales", "Partner & sales manager (accountant channel) gross, full-time equivalent", 2000, "EUR/month", EUR,
     "ASSUMPTION (no sourced salary benchmark). Paid pro rata to in_SalesFTE while part-time."),
    ("SalesStart", "Partner & sales manager start month", 3, "model month", NUM, "ASSUMPTION."),
    ("SalesFTE", "Partner & sales manager share of full time until full-time start", 0.5, "FTE", PCT,
     "Half time in year 1 to fit the EUR 150k round. Partner signings scale with it (new partners/month x FTE): fewer sales hours = fewer partners. ASSUMPTION."),
    ("SalesFullStart", "Partner & sales manager full time from", 13, "model month", NUM, "Jan 2028, after the first year of channel data. ASSUMPTION."),
    ("SalSupport", "Customer onboarding & support specialist (HR/payroll background) gross, full-time equivalent", 1400, "EUR/month", EUR,
     "ASSUMPTION. Paid pro rata to in_SupportFTE while part-time."),
    ("SupportStart", "Onboarding & support specialist start month", 7, "model month", NUM,
     "Founders onboard the pilot companies (Apr 2027); the specialist starts part-time around 30 paying customers. ASSUMPTION."),
    ("SupportFTE", "Onboarding & support specialist share of full time until full-time start", 0.5, "FTE", PCT, "ASSUMPTION."),
    ("SupportFullStart", "Onboarding & support specialist full time from", 25, "model month", NUM,
     "With the seed (same month as the base-case developer hire). ASSUMPTION."),
    ("SalDev", "Full-stack developer gross", 2500, "EUR/month", EUR, "ASSUMPTION. Start month is a scenario driver (in_DevStart)."),
    ("#", "Operating costs (EUR)"),
    ("InfraFixed", "Hosting, database, e-mail, monitoring, backups (fixed)", 350, "EUR/month", EUR,
     "Founders' internal cost estimate ~EUR 2,000/yr minimal; raised for a separate prod project + staging."),
    ("InfraVar", "Variable infra per paying customer (storage, push, AI usage)", 1.5, "EUR/customer/mo", EUR2, "ASSUMPTION."),
    ("Tools", "Development tools and AI subscriptions", 250, "EUR/month", EUR, "Founders' internal cost estimate (EUR 2,400/yr), rounded up."),
    ("PayFee", "Payment processing fees", 0.02, "% of revenue", PCT, "No billing provider integrated yet (invoicing is manual today). ASSUMPTION for card/online payment."),
    ("MktY1", "Marketing budget, year 1", 1500, "EUR/month", EUR, "ASSUMPTION."),
    ("MktY2", "Marketing budget, year 2", 2500, "EUR/month", EUR, "ASSUMPTION."),
    ("MktY3", "Marketing budget, year 3", 3500, "EUR/month", EUR, "ASSUMPTION."),
    ("Travel", "Travel to accounting practices, accountant events", 300, "EUR/month", EUR,
     "Founders' internal cost estimate (accountant sales), trimmed to fit the EUR 150k round: local Timiș/West Region practices first, few paid events."),
    ("Office", "Coworking desk(s) in Timiș county", 200, "EUR/month", EUR, "VV requires an operating site in the West Region (vestventures.vc/en/programs-terms). ASSUMPTION."),
    ("Insurance", "Professional liability insurance", 100, "EUR/month", EUR, "Founders' internal cost estimate (EUR 1,200/yr)."),
    ("Admin", "Bookkeeping, bank, admin", 200, "EUR/month", EUR, "ASSUMPTION."),
    ("Contingency", "Contingency on all costs", 0.05, "%", PCT, "ASSUMPTION."),
    ("#", "One-off legal, validation and compliance (EUR, spread evenly from start to end month)"),
    ("LegalAmt", "Legal: round documents, T&C, DPA, IP assignment to the SRL", 5000, "EUR", EUR, "Founders' internal cost estimate (EUR 3,500 minimal) + round paperwork."),
    ("LegalStart", "  start month", 1, "model month", NUM, ""),
    ("LegalEnd", "  end month", 2, "model month", NUM, ""),
    ("PayrollValAmt", "Payroll legal values validated by accountant / lawyer", 8000, "EUR", EUR,
     "Founders' internal cost estimate ('~EUR 8,000 of expertise'). Every payroll legal value is still marked to confirm."),
    ("PayrollValStart", "  start month", 2, "model month", NUM, ""),
    ("PayrollValEnd", "  end month", 5, "model month", NUM, ""),
    ("A11yAmt", "Accessibility audit and fixes (EN 301 549)", 3000, "EUR", EUR, "Required for software built inside the VV programme (vestventures.vc/en/programs-terms). ASSUMPTION on cost."),
    ("A11yStart", "  start month", 5, "model month", NUM, ""),
    ("A11yEnd", "  end month", 6, "model month", NUM, ""),
    ("PentestAmt", "External penetration test", 6000, "EUR", EUR, "Founders' internal cost estimate ('~EUR 6,000')."),
    ("PentestStart", "  start month", 12, "model month", NUM, "Dec 2027: before the year-2 scale-up and before seed due diligence."),
    ("PentestEnd", "  end month", 12, "model month", NUM, ""),
    ("#", "Tax"),
    ("TaxRegime", "Tax regime (1 = micro-enterprise on revenue, 2 = profit tax)", 1, "1 / 2", NUM,
     "[TO CONFIRM] with accountant: micro-enterprise eligibility (revenue threshold, CAEN, shareholder tests) for the SRL."),
    ("MicroRate", "Micro-enterprise tax on revenue", 0.01, "% of revenue", PCT, "[TO CONFIRM] 1% vs 3% depending on CAEN code / revenue."),
    ("ProfitTax", "Profit tax (if regime 2)", 0.16, "% of profit", PCT, "Standard Romanian rate; no loss carry-forward modelled (prudent)."),
    ("#", "Analysis settings"),
    ("UEYear", "Model year used for unit economics", 2, "year", NUM, "Year 2 = first full year with the whole team; year 1 is distorted by pilot conversion."),
]


# --------------------------------------------------------------------------------------
# Builders
# --------------------------------------------------------------------------------------

class Model:
    def __init__(self, scenario: int, snapshot: dict | None):
        self.wb = Workbook()
        self.scenario = scenario
        self.snapshot = snapshot
        self.names: dict[str, str] = {}  # in_X -> absolute ref
        self.rows: dict[str, dict[str, int]] = {}

    # -- names ------------------------------------------------------------------------
    def name(self, key: str, sheet: str, cell: str):
        col_, row_ = "".join(ch for ch in cell if ch.isalpha()), "".join(ch for ch in cell if ch.isdigit())
        target = f"{ref(sheet, '$' + col_ + '$' + row_)}"
        nm = f"in_{key}"
        self.wb.defined_names[nm] = DefinedName(nm, attr_text=target)
        self.names[nm] = target

    # -- generic monthly sheet helpers -----------------------------------------------
    def month_header(self, ws, title: str, subtitle: str):
        ws["A1"] = title
        ws["A1"].font = TITLE
        ws["A2"] = subtitle
        ws["A2"].font = SUB
        ws.column_dimensions["A"].width = 52
        ws.column_dimensions["B"].width = 16
        ws["A3"], ws["A4"], ws["A5"] = "Month #", "Month", "Model year"
        for a in ("A3", "A4", "A5"):
            ws[a].font = BOLD
        for m in range(1, N_MONTHS + 1):
            c = col(m)
            ws.column_dimensions[c].width = 11
            if ws.title == "Revenue":
                ws[f"{c}3"] = 1 if m == 1 else f"={col(m - 1)}3+1"
                ws[f"{c}4"] = f"=DATE(YEAR(in_StartDate),MONTH(in_StartDate)+{c}3-1,1)"
                ws[f"{c}5"] = f"=INT(({c}3-1)/12)+1"
                for r in (3, 4, 5):
                    ws[f"{c}{r}"].font = BOLD
            else:
                for r in (3, 4, 5):
                    ws[f"{c}{r}"] = f"=Revenue!{c}{r}"
                    ws[f"{c}{r}"].font = Font(name=FONT, color="008000", bold=True)
            ws[f"{c}4"].number_format = "mmm yy"
            for r in (3, 4, 5):
                ws[f"{c}{r}"].fill = SEC_FILL
                ws[f"{c}{r}"].alignment = Alignment(horizontal="center")
        for r in (3, 4, 5):
            ws[f"A{r}"].fill = SEC_FILL
            ws[f"B{r}"].fill = SEC_FILL
        ws.freeze_panes = "C6"

    def mrow(self, ws, row: int, key: str, label: str, unit: str, fn, fmt=NUM, font=BLACK, bold=False, border=False):
        """Write a monthly row: fn(m, c) returns the formula for month m in column c."""
        self.rows.setdefault(ws.title, {})[key] = row
        ws[f"A{row}"] = label
        ws[f"B{row}"] = unit
        ws[f"B{row}"].font = SUB
        if bold:
            ws[f"A{row}"].font = BOLD
        for m in range(1, N_MONTHS + 1):
            c = col(m)
            cell = ws[f"{c}{row}"]
            cell.value = fn(m, c)
            cell.number_format = fmt
            cell.font = Font(name=FONT, color=font.color, bold=bold)
            if border:
                cell.border = TOP
        return row

    def section(self, ws, row: int, text: str, width: int = LAST_COL):
        ws[f"A{row}"] = text
        ws[f"A{row}"].font = BOLD
        for ci in range(1, width + 1):
            ws.cell(row=row, column=ci).fill = SEC_FILL

    def r(self, sheet: str, key: str) -> int:
        return self.rows[sheet][key]

    # ---------------------------------------------------------------------------------
    def build(self):
        wb = self.wb
        self.cover = wb.active
        self.cover.title = "Cover"
        for t in ["Assumptions", "Revenue", "Costs", "P&L & Cash", "Unit economics", "Scenarios", "Use of funds", "Cap table", "Sources"]:
            wb.create_sheet(t)
        self.build_scenarios()
        self.build_assumptions()
        self.build_revenue()
        self.build_costs()
        self.build_pnl()
        self.build_unit_economics()
        self.build_use_of_funds()
        self.build_cap_table()
        self.build_sources()
        self.build_cover()
        self.fill_snapshot()
        for ws in wb.worksheets:
            style_all(ws)
            ws.sheet_view.showGridLines = False
        wb.calculation.fullCalcOnLoad = True
        return wb

    # -- Scenarios ---------------------------------------------------------------------
    def build_scenarios(self):
        ws = self.wb["Scenarios"]
        ws["A1"] = "Scenarios - key growth drivers"
        ws["A1"].font = TITLE
        ws["A2"] = ("Blue = input. The scenario in use is chosen on Cover!C8 (1 Base, 2 Conservative, 3 Upside); "
                    "Assumptions picks the matching column with INDEX. All values are ASSUMPTIONS / PROJECTIONS.")
        ws["A2"].font = SUB
        heads = ["Driver", "Unit", "Base", "Conservative", "Upside", "Note / source"]
        for i, h in enumerate(heads, 1):
            c = ws.cell(row=4, column=i, value=h)
            c.font, c.fill = HDR, HDR_FILL
        ws.column_dimensions["A"].width = 48
        ws.column_dimensions["B"].width = 12
        for cc in "CDE":
            ws.column_dimensions[cc].width = 14
        ws.column_dimensions["F"].width = 100
        self.scen_rows = {}
        for i, (key, label, b, c_, u, unit, fmt, note) in enumerate(SCENARIO_DRIVERS):
            row = 5 + i
            self.scen_rows[key] = row
            ws[f"A{row}"], ws[f"B{row}"] = label, unit
            for colL, v in zip("CDE", (b, c_, u)):
                ws[f"{colL}{row}"] = v
                ws[f"{colL}{row}"].font = BLUE
                ws[f"{colL}{row}"].number_format = fmt
            ws[f"F{row}"] = note
            ws[f"F{row}"].font = SUB
        self.scen_last = 5 + len(SCENARIO_DRIVERS) - 1
        self.snap_row = self.scen_last + 3

    def fill_snapshot(self):
        ws = self.wb["Scenarios"]
        row = self.snap_row
        self.section(ws, row, "Scenario results snapshot (PROJECTION)", 6)
        ws[f"A{row + 1}"] = ("Values below were produced by recalculating this workbook once per scenario (model/build_model.py + "
                             "model/recalc.sh) and pasted as values so that a reader can compare scenarios without switching. "
                             "Live values for the selected scenario: 'P&L & Cash' and Cover.")
        ws[f"A{row + 1}"].font = SUB
        heads = ["Output", "Unit", "Base", "Conservative", "Upside"]
        for i, h in enumerate(heads, 1):
            c = ws.cell(row=row + 2, column=i, value=h)
            c.font, c.fill = HDR, HDR_FILL
        items = [
            ("customers_m12", "Paying customers, month 12 (Dec 2027)", "firms", NUM),
            ("mrr_m12_eur", "MRR, month 12", "EUR", EUR),
            ("customers_m24", "Paying customers, month 24 (Dec 2028)", "firms", NUM),
            ("mrr_m24_eur", "MRR, month 24", "EUR", EUR),
            ("customers_m36", "Paying customers, month 36 (Dec 2029)", "firms", NUM),
            ("mrr_m36_eur", "MRR, month 36", "EUR", EUR),
            ("arr_m36_eur", "ARR, month 36", "EUR", EUR),
            ("breakeven_month", "Break-even month (first month net cash flow >= 0)", "model month", NUM),
            ("cash_low_point_eur", "Cash low point", "EUR", EUR),
            ("cash_m36_eur", "Cash at month 36", "EUR", EUR),
            ("months_to_30", "Month reaching 30 paying customers", "model month", NUM),
            ("months_to_100", "Month reaching 100 paying customers", "model month", NUM),
            ("months_to_300", "Month reaching 300 paying customers", "model month", NUM),
        ]
        for i, (k, label, unit, fmt) in enumerate(items):
            rr = row + 3 + i
            ws[f"A{rr}"], ws[f"B{rr}"] = label, unit
            for colL, sname in zip("CDE", SCEN_NAMES):
                v = (self.snapshot or {}).get(sname, {}).get(k, "run recalc.sh")
                ws[f"{colL}{rr}"] = v
                ws[f"{colL}{rr}"].number_format = fmt
                ws[f"{colL}{rr}"].font = SUB if isinstance(v, str) else BLACK

    # -- Assumptions -------------------------------------------------------------------
    def build_assumptions(self):
        ws = self.wb["Assumptions"]
        ws["A1"] = "Assumptions - every driver of the model"
        ws["A1"].font = TITLE
        ws["A2"] = ("Blue = hard-coded input (edit here). Green = picked from 'Scenarios' for the scenario chosen on Cover. "
                    "Black = calculated. Yellow = key lever. Nothing on other sheets is hard-coded.")
        ws["A2"].font = SUB
        for i, h in enumerate(["Driver", "Value", "Unit", "Source / note"], 1):
            c = ws.cell(row=4, column=i, value=h)
            c.font, c.fill = HDR, HDR_FILL
        ws.column_dimensions["A"].width = 70
        ws.column_dimensions["B"].width = 16
        ws.column_dimensions["C"].width = 16
        ws.column_dimensions["D"].width = 110
        row = 5
        key_levers = {"VVTicket", "Angels", "PreMoney", "ESOP"}
        for item in INPUTS:
            if item[0] == "#":
                row += 1
                self.section(ws, row, item[1], 4)
                row += 1
                continue
            key, label, val, unit, fmt, note = item
            ws[f"A{row}"], ws[f"C{row}"], ws[f"D{row}"] = label, unit, note
            if key == "StartDate":
                import datetime as dt
                val = dt.datetime.strptime(val, "%Y-%m-%d")
            ws[f"B{row}"] = val
            ws[f"B{row}"].font = BLUE
            ws[f"B{row}"].number_format = fmt
            if key in key_levers:
                ws[f"B{row}"].fill = YELLOW
            ws[f"D{row}"].font = SUB
            self.name(key, "Assumptions", f"B{row}")
            row += 1
            if key == "OpeningCash":
                pass
            if key == "Angels":
                ws[f"A{row}"], ws[f"C{row}"] = "Round size (VV ticket + angels)", "EUR"
                ws[f"B{row}"] = "=in_VVTicket+in_Angels"
                ws[f"B{row}"].number_format = EUR
                ws[f"B{row}"].font = BOLD
                ws[f"D{row}"] = "Calculated. Target range given by the founders: EUR 150k-300k pre-seed."
                ws[f"D{row}"].font = SUB
                self.name("Round", "Assumptions", f"B{row}")
                row += 1
            if key == "MixAll":
                ws[f"A{row}"], ws[f"C{row}"] = "Check: plan mix total", "%"
                ws[f"B{row}"] = "=" + "+".join(f"in_Mix{k}" for k, *_ in PRICES_RON)
                ws[f"B{row}"].number_format = PCT
                ws[f"D{row}"] = "Must equal 100%."
                ws[f"D{row}"].font = SUB
                self.name("MixCheck", "Assumptions", f"B{row}")
                row += 1
                ws[f"A{row}"], ws[f"C{row}"] = "Blended list price per paying customer (year 1)", "RON/month"
                ws[f"B{row}"] = "=" + "+".join(f"in_Price{k}*in_Mix{k}" for k, *_ in PRICES_RON)
                ws[f"B{row}"].number_format = '#,##0.0" RON"'
                ws[f"B{row}"].font = BOLD
                ws[f"D{row}"] = ("Calculated = sum(price x mix). Firms above 20 employees have no published price "
                                 "-> modelled at list price (understates revenue from larger firms).")
                ws[f"D{row}"].font = SUB
                self.name("BlendedPriceRON", "Assumptions", f"B{row}")
                row += 1
                ws[f"A{row}"], ws[f"C{row}"] = "Blended ARPA, year 1", "EUR/month"
                ws[f"B{row}"] = "=in_BlendedPriceRON/in_FX"
                ws[f"B{row}"].number_format = EUR2
                ws[f"B{row}"].font = BOLD
                self.name("ARPA1", "Assumptions", f"B{row}")
                row += 1
        # scenario drivers
        row += 1
        self.section(ws, row, "Scenario drivers (selected scenario, from 'Scenarios')", 4)
        row += 1
        ws[f"A{row}"] = "Scenario in use"
        ws[f"B{row}"] = '=INDEX(Scenarios!$C$4:$E$4,1,ScenarioIdx)'
        ws[f"B{row}"].font = GREEN
        row += 1
        for key, label, *_rest in SCENARIO_DRIVERS:
            unit, fmt, note = _rest[3], _rest[4], _rest[5]
            sr = self.scen_rows[key]
            ws[f"A{row}"], ws[f"C{row}"] = label, unit
            ws[f"B{row}"] = f"=INDEX(Scenarios!$C${sr}:$E${sr},1,ScenarioIdx)"
            ws[f"B{row}"].font = GREEN
            ws[f"B{row}"].number_format = fmt
            ws[f"D{row}"] = note
            ws[f"D{row}"].font = SUB
            self.name(key, "Assumptions", f"B{row}")
            row += 1

    # -- Revenue -----------------------------------------------------------------------
    def build_revenue(self):
        ws = self.wb["Revenue"]
        S = "Revenue"
        self.month_header(ws, "Revenue - customers, MRR, ARR (PROJECTION, selected scenario)",
                          "Month 1 = Jan 2027 = assumed close of the round. Zero paying customers on 6 Oct 2026. "
                          "Customers are expected values (fractions allowed). EUR at in_FX RON/EUR.")
        prev = lambda c, row, m: f"{col(m - 1)}{row}" if m > 1 else "0"
        r = 7
        self.section(ws, r, "Acquisition channels")
        self.mrow(ws, 8, "pilot", "Pilot companies converting to paid", "firms",
                  lambda m, c: f"=IF({c}$3=in_PilotConvMonth,in_PilotPractices*in_PilotFirmsPerPractice*in_PilotConv,0)", NUM1)
        self.mrow(ws, 9, "newpartners", "New accountant partners signed (x partner manager FTE)", "practices",
                  lambda m, c: (f"=IF({c}$3>=in_PartnerStart,CHOOSE({c}$5,in_NewPartnersY1,in_NewPartnersY2,in_NewPartnersY3)"
                                f"*IF({c}$3>=in_SalesFullStart,1,in_SalesFTE),0)"), NUM1)
        self.mrow(ws, 10, "partners", "Active accountant partners (end of month)", "practices",
                  lambda m, c: f"={prev(c, 10, m)}+{c}9+IF({c}$3=in_PilotConvMonth,in_PilotPractices,0)", NUM1)
        self.mrow(ws, 11, "referred", "Firms referred by partners", "firms",
                  lambda m, c: f"={prev(c, 10, m)}*in_RefRate", NUM1)
        self.mrow(ws, 12, "newpartnercust", "New paying customers - partner referrals (after free month)", "firms",
                  lambda m, c: f"={prev(c, 11, m)}*in_RefConv", NUM1)
        self.mrow(ws, 13, "organic", "Organic / SEO / direct trials", "trials",
                  lambda m, c: f"=CHOOSE({c}$5,in_OrganicY1,in_OrganicY2,in_OrganicY3)", NUM1)
        self.mrow(ws, 14, "paidtrials", "Paid-marketing trials", "trials",
                  lambda m, c: f"=Costs!{c}{{MKT}}*in_PaidShare/in_CostPerTrial", NUM1, GREEN)
        self.mrow(ws, 15, "newdirect", "New paying customers - direct (after free month)", "firms",
                  lambda m, c: f"=({prev(c, 13, m)}+{prev(c, 14, m)})*in_TrialConv", NUM1)
        self.section(ws, 17, "Customers")
        self.mrow(ws, 18, "new", "New paying customers - total", "firms", lambda m, c: f"={c}8+{c}12+{c}15", NUM1, bold=True)
        self.mrow(ws, 19, "newaccountant", "  of which via accountants (pilot + partners)", "firms", lambda m, c: f"={c}8+{c}12", NUM1)
        self.mrow(ws, 20, "churned", "Churned customers", "firms", lambda m, c: f"={prev(c, 21, m)}*in_Churn", NUM1)
        self.mrow(ws, 21, "customers", "Paying customers (end of month)", "firms",
                  lambda m, c: f"={prev(c, 21, m)}+{c}18-{c}20", NUM1, bold=True, border=True)
        self.section(ws, 23, "Revenue")
        self.mrow(ws, 24, "priceRON", "Blended list price", "RON/month", lambda m, c: f"=in_BlendedPriceRON*(1+in_PriceInc)^({c}$5-1)", '#,##0.0')
        self.mrow(ws, 25, "arpa", "ARPA", "EUR/month", lambda m, c: f"={c}24/in_FX", EUR2)
        self.mrow(ws, 26, "mrr", "MRR", "EUR", lambda m, c: f"={c}21*{c}25", EUR, bold=True)
        self.mrow(ws, 27, "arr", "ARR (MRR x 12)", "EUR", lambda m, c: f"={c}26*12", EUR)
        self.mrow(ws, 28, "revenue", "Revenue recognised (billed monthly)", "EUR", lambda m, c: f"={c}26", EUR, bold=True, border=True)
        # thresholds
        self.section(ws, 30, "Milestones (selected scenario)", 6)
        crow = rng(S, 21)
        mrow_ = rng(S, 3)
        drow = rng(S, 4)
        self.thr_rows = {}
        for i, n in enumerate((30, 100, 300)):
            rr = 31 + i
            ws[f"A{rr}"] = f"First month with >= {n} paying customers"
            ws[f"B{rr}"] = n
            ws[f"B{rr}"].font = BLUE
            ws[f"C{rr}"] = f'=IF(MAX({crow})<B{rr},"not within 36 months",COUNTIF({crow},"<"&B{rr})+1)'
            ws[f"D{rr}"] = f'=IF(ISNUMBER(C{rr}),INDEX({drow},1,C{rr}),"-")'
            ws[f"D{rr}"].number_format = "mmm yyyy"
            ws[f"E{rr}"] = "month #, date (COUNTIF works because the customer curve is non-decreasing)"
            ws[f"E{rr}"].font = SUB
            self.thr_rows[n] = rr
        self.name("M30", S, "C31")
        self.name("M100", S, "C32")
        self.name("M300", S, "C33")

    # -- Costs -------------------------------------------------------------------------
    def build_costs(self):
        ws = self.wb["Costs"]
        self.month_header(ws, "Costs - headcount plan and operating expenses (EUR, PROJECTION)",
                          "Salaries = gross x (1 + CAM) x (1 + raise)^(year-1). One-off items are spread evenly between their start and end month.")
        rev = lambda c, key: f"Revenue!{c}{self.r('Revenue', key)}"
        self.section(ws, 7, "Headcount (employer cost)")
        sal = lambda gross, start: (lambda m, c: f"=IF({c}$3>={start},{gross}*(1+in_CAM)*(1+in_Raise)^({c}$5-1),0)")
        # part-time until full_start: gross x FTE share
        salpt = lambda gross, start, fte, full_start: (lambda m, c: (f"=IF({c}$3>={start},{gross}*IF({c}$3>={full_start},1,{fte})"
                                                                     f"*(1+in_CAM)*(1+in_Raise)^({c}$5-1),0)"))
        fte_of = lambda c, start, fte, full_start: f"IF({c}$3>={start},IF({c}$3>={full_start},1,{fte}),0)"
        self.mrow(ws, 8, "founder1", "Co-founder - Miroslav Maletici (role [TO CONFIRM])", "EUR", sal("in_SalFounder", 1), EUR)
        self.mrow(ws, 9, "founder2", "Co-founder - Răzvan Pervulescu (role [TO CONFIRM])", "EUR", sal("in_SalFounder", 1), EUR)
        self.mrow(ws, 10, "sales", "Partner & sales manager (accountant channel; part-time first)", "EUR",
                  salpt("in_SalSales", "in_SalesStart", "in_SalesFTE", "in_SalesFullStart"), EUR)
        self.mrow(ws, 11, "support", "Customer onboarding & support specialist (part-time first)", "EUR",
                  salpt("in_SalSupport", "in_SupportStart", "in_SupportFTE", "in_SupportFullStart"), EUR)
        self.mrow(ws, 12, "dev", "Full-stack developer", "EUR", sal("in_SalDev", "in_DevStart"), EUR)
        self.mrow(ws, 13, "people", "Total headcount cost", "EUR", lambda m, c: f"=SUM({c}8:{c}12)", EUR, bold=True, border=True)
        self.mrow(ws, 14, "fte", "Headcount (FTE)", "FTE",
                  lambda m, c: (f"=2+{fte_of(c, 'in_SalesStart', 'in_SalesFTE', 'in_SalesFullStart')}"
                                f"+{fte_of(c, 'in_SupportStart', 'in_SupportFTE', 'in_SupportFullStart')}+IF({c}$3>=in_DevStart,1,0)"), NUM1)
        self.section(ws, 16, "Operating expenses")
        self.mrow(ws, 17, "infrafix", "Hosting & infrastructure (fixed)", "EUR", lambda m, c: "=in_InfraFixed", EUR)
        self.mrow(ws, 18, "infravar", "Infrastructure (variable, per customer)", "EUR", lambda m, c: f"={rev(c, 'customers')}*in_InfraVar", EUR, GREEN)
        self.mrow(ws, 19, "tools", "Development tools & AI subscriptions", "EUR", lambda m, c: "=in_Tools", EUR)
        self.mrow(ws, 20, "payfee", "Payment processing fees", "EUR", lambda m, c: f"={rev(c, 'revenue')}*in_PayFee", EUR, GREEN)
        self.mrow(ws, 21, "mkt", "Marketing (content, SEO, ads, materials)", "EUR",
                  lambda m, c: f"=CHOOSE({c}$5,in_MktY1,in_MktY2,in_MktY3)", EUR)
        self.mrow(ws, 22, "travel", "Travel & accountant events", "EUR", lambda m, c: "=in_Travel", EUR)
        mrow_rev, acc_rev = rng("Revenue", 3), rng("Revenue", self.r("Revenue", "newaccountant"))
        self.mrow(ws, 23, "comm", "Accountant commissions (20% x 6 months)", "EUR",
                  lambda m, c: (f"=in_CommRate*{rev(c, 'arpa')}*SUMPRODUCT(({mrow_rev}>{c}$3-in_CommMonths)*"
                                f"({mrow_rev}<={c}$3)*{acc_rev})"), EUR, GREEN)
        self.mrow(ws, 24, "office", "Coworking (Timiș)", "EUR", lambda m, c: "=in_Office", EUR)
        self.mrow(ws, 25, "insurance", "Professional liability insurance", "EUR", lambda m, c: "=in_Insurance", EUR)
        self.mrow(ws, 26, "admin", "Bookkeeping, bank, admin", "EUR", lambda m, c: "=in_Admin", EUR)
        one = lambda k: (lambda m, c: f"=IF(AND({c}$3>=in_{k}Start,{c}$3<=in_{k}End),in_{k}Amt/(in_{k}End-in_{k}Start+1),0)")
        self.mrow(ws, 27, "legal", "Legal: round, T&C, DPA, IP assignment (one-off)", "EUR", one("Legal"), EUR)
        self.mrow(ws, 28, "payrollval", "Payroll legal-value validation (one-off)", "EUR", one("PayrollVal"), EUR)
        self.mrow(ws, 29, "a11y", "Accessibility audit EN 301 549 (one-off)", "EUR", one("A11y"), EUR)
        self.mrow(ws, 30, "pentest", "Penetration test (one-off)", "EUR", one("Pentest"), EUR)
        self.mrow(ws, 31, "opex", "Total operating expenses", "EUR", lambda m, c: f"=SUM({c}17:{c}30)", EUR, bold=True, border=True)
        self.mrow(ws, 32, "contingency", "Contingency", "EUR", lambda m, c: f"=({c}13+{c}31)*in_Contingency", EUR)
        self.mrow(ws, 34, "total", "TOTAL COSTS", "EUR", lambda m, c: f"={c}13+{c}31+{c}32", EUR, bold=True, border=True)
        # patch the forward reference from Revenue (paid trials) now that the marketing row is known
        wsr = self.wb["Revenue"]
        for m in range(1, N_MONTHS + 1):
            c = col(m)
            wsr[f"{c}14"].value = wsr[f"{c}14"].value.replace("{MKT}", str(self.r("Costs", "mkt")))

    # -- P&L & Cash --------------------------------------------------------------------
    def build_pnl(self):
        ws = self.wb["P&L & Cash"]
        S = "P&L & Cash"
        self.month_header(ws, "P&L and cash (EUR, PROJECTION, selected scenario)",
                          "Gross profit treats hosting, payment fees and the onboarding/support salary as cost of revenue. "
                          "Break-even = first month with net cash flow >= 0.")
        cst = lambda c, key: f"Costs!{c}{self.r('Costs', key)}"
        self.section(ws, 7, "Profit and loss")
        self.mrow(ws, 8, "rev", "Revenue", "EUR", lambda m, c: f"=Revenue!{c}{self.r('Revenue', 'revenue')}", EUR, GREEN, bold=True)
        self.mrow(ws, 9, "cogs", "Cost of revenue (hosting, fees, onboarding & support)", "EUR",
                  lambda m, c: f"={cst(c, 'infrafix')}+{cst(c, 'infravar')}+{cst(c, 'payfee')}+{cst(c, 'support')}", EUR, GREEN)
        self.mrow(ws, 10, "gp", "Gross profit", "EUR", lambda m, c: f"={c}8-{c}9", EUR, bold=True, border=True)
        self.mrow(ws, 11, "gm", "Gross margin", "%", lambda m, c: f"=IF({c}8>0,{c}10/{c}8,0)", PCT)
        self.mrow(ws, 12, "rnd", "R&D (co-founders, developer, tools)", "EUR",
                  lambda m, c: f"={cst(c, 'founder1')}+{cst(c, 'founder2')}+{cst(c, 'dev')}+{cst(c, 'tools')}", EUR, GREEN)
        self.mrow(ws, 13, "snm", "Sales & marketing (partner manager, marketing, travel, commissions)", "EUR",
                  lambda m, c: f"={cst(c, 'sales')}+{cst(c, 'mkt')}+{cst(c, 'travel')}+{cst(c, 'comm')}", EUR, GREEN)
        self.mrow(ws, 14, "gna", "G&A (legal, validation, compliance, office, insurance, admin, contingency)", "EUR",
                  lambda m, c: (f"={cst(c, 'office')}+{cst(c, 'insurance')}+{cst(c, 'admin')}+{cst(c, 'legal')}+"
                                f"{cst(c, 'payrollval')}+{cst(c, 'a11y')}+{cst(c, 'pentest')}+{cst(c, 'contingency')}"), EUR, GREEN)
        self.mrow(ws, 15, "ebitda", "EBITDA", "EUR", lambda m, c: f"={c}10-{c}12-{c}13-{c}14", EUR, bold=True, border=True)
        self.mrow(ws, 16, "check", "Check: EBITDA = revenue - total costs (should be 0)", "EUR",
                  lambda m, c: f"=ROUND({c}15-({c}8-{cst(c, 'total')}),2)", EUR)
        self.mrow(ws, 17, "tax", "Tax (micro on revenue, or profit tax)", "EUR",
                  lambda m, c: f"=IF(in_TaxRegime=1,in_MicroRate*{c}8,MAX(0,in_ProfitTax*{c}15))", EUR)
        self.mrow(ws, 18, "net", "Net result = operating cash flow", "EUR", lambda m, c: f"={c}15-{c}17", EUR, bold=True, border=True)
        self.section(ws, 20, "Cash")
        prev = lambda row, m: f"{col(m - 1)}{row}" if m > 1 else None
        self.mrow(ws, 21, "open", "Opening cash", "EUR", lambda m, c: f"={prev(23, m)}" if m > 1 else "=in_OpeningCash", EUR)
        self.mrow(ws, 22, "funding", "Funding received (pre-seed round)", "EUR", lambda m, c: f"=IF({c}$3=in_FundingMonth,in_Round,0)", EUR)
        self.mrow(ws, 23, "cash", "Closing cash", "EUR", lambda m, c: f"={c}21+{c}22+{c}18", EUR, bold=True, border=True)
        self.mrow(ws, 24, "cumcost", "Cumulative total costs (for zero-revenue runway)", "EUR",
                  lambda m, c: f"={prev(24, m)}+{cst(c, 'total')}" if m > 1 else f"={cst(c, 'total')}", EUR)
        self.mrow(ws, 25, "cumnet", "Cumulative operating cash flow", "EUR",
                  lambda m, c: f"={prev(25, m)}+{c}18" if m > 1 else f"={c}18", EUR)
        self.mrow(ws, 26, "negflag", "Flag: cash below zero", "1/0", lambda m, c: f"=IF({c}23<0,1,0)", NUM)
        self.mrow(ws, 27, "beflag", "Flag: net cash flow >= 0", "1/0", lambda m, c: f"=IF({c}18>=0,1,0)", NUM)
        self.mrow(ws, 28, "cashA", "Closing cash - option A (VV ticket + minimum private match only)", "EUR",
                  lambda m, c: f"={c}23-(in_Round-in_OptionA)*IF({c}$3>=in_FundingMonth,1,0)", EUR)
        self.mrow(ws, 29, "negflagA", "Flag: option A cash below zero", "1/0", lambda m, c: f"=IF({c}28<0,1,0)", NUM)
        self.mrow(ws, 30, "runmin", "Running minimum of cumulative operating cash flow (for round-size tests)", "EUR",
                  lambda m, c: f"=MIN({prev(30, m)},{c}25)" if m > 1 else f"={c}25", EUR)

        # summary block
        cash, cumcost = rng(S, 23), rng(S, 24)
        neg, be, negA, cashA = rng(S, 26), rng(S, 27), rng(S, 29), rng(S, 28)
        dates = rng("Revenue", 4)
        self.section(ws, 31, "Summary (selected scenario)", 6)
        rows = [
            ("OptionA", "Option A round: VV ticket + minimum private co-investment", "=in_VVTicket*(1+in_MinPrivatePct)", EUR,
             "VV ticket alone is not allowed: VV requires >=10% independent private money."),
            ("RunwayZeroRev", "Runway with ZERO revenue (target round)", f'=COUNTIF({cumcost},"<="&(in_Round+in_OpeningCash))', MON,
             "Months the round pays every planned cost even if no customer ever pays. Basis of the round size."),
            ("RunwayZeroRevA", "Runway with ZERO revenue (option A)", f'=COUNTIF({cumcost},"<="&(in_OptionA+in_OpeningCash))', MON, ""),
            ("RunwayPlan", "Runway on plan (months until cash < 0)", f'=IF(MAX({neg})=0,"> 36",MATCH(1,{neg},0)-1)', MON,
             "'> 36' = cash never goes negative within the model horizon."),
            ("RunwayPlanA", "Runway on plan, option A", f'=IF(MAX({negA})=0,"> 36",MATCH(1,{negA},0)-1)', MON, ""),
            ("Breakeven", "Break-even month (first month net cash flow >= 0)", f'=IF(MAX({be})=0,"not within 36 months",MATCH(1,{be},0))', NUM,
             "model month #; date in next cell"),
            ("CashLow", "Cash low point", f"=MIN({cash})", EUR, ""),
            ("CashLowMonth", "Cash low point month", f"=MATCH(in_CashLow,{cash},0)", NUM, ""),
            ("CashLowA", "Cash low point, option A", f"=MIN({cashA})", EUR, ""),
            ("Cash36", "Cash at month 36", f"={CN}23", EUR, ""),
            ("Rev12", "Revenue, months 1-12", f"=SUMIF({rng('Revenue', 5)},1,{rng(S, 8)})", EUR, ""),
            ("Rev24", "Revenue, months 13-24", f"=SUMIF({rng('Revenue', 5)},2,{rng(S, 8)})", EUR, ""),
            ("Rev36", "Revenue, months 25-36", f"=SUMIF({rng('Revenue', 5)},3,{rng(S, 8)})", EUR, ""),
            ("Cost12", "Total costs, months 1-12", f"=SUMIF({rng('Revenue', 5)},1,{rng('Costs', 34)})", EUR, ""),
            ("Cost24", "Total costs, months 13-24", f"=SUMIF({rng('Revenue', 5)},2,{rng('Costs', 34)})", EUR, ""),
            ("Cost36", "Total costs, months 25-36", f"=SUMIF({rng('Revenue', 5)},3,{rng('Costs', 34)})", EUR, ""),
            ("MaxCheck", "Integrity check: max |EBITDA check| (must be 0)", f"=MAX({rng(S, 16)})-MIN({rng(S, 16)})", EUR2, ""),
        ]
        for i, (key, label, f, fmt, note) in enumerate(rows):
            rr = 32 + i
            ws[f"A{rr}"], ws[f"C{rr}"], ws[f"E{rr}"] = label, f, note
            ws[f"C{rr}"].number_format = fmt
            ws[f"C{rr}"].font = BOLD
            ws[f"E{rr}"].font = SUB
            self.name(key, S, f"C{rr}")
            if key == "Breakeven":
                ws[f"D{rr}"] = f'=IF(ISNUMBER(C{rr}),INDEX({dates},1,C{rr}),"-")'
                ws[f"D{rr}"].number_format = "mmm yyyy"
            if key == "CashLowMonth":
                ws[f"D{rr}"] = f"=INDEX({dates},1,C{rr})"
                ws[f"D{rr}"].number_format = "mmm yyyy"

    # -- Unit economics ----------------------------------------------------------------
    def build_unit_economics(self):
        ws = self.wb["Unit economics"]
        ws["A1"] = "Unit economics (PROJECTION, selected scenario, model year in_UEYear)"
        ws["A1"].font = TITLE
        ws["A2"] = "Computed over one model year (default year 2 = Jan-Dec 2028). Benchmarks: ChartMogul (churn) and Aleph (CAC payback, secondary); URLs on sheet Sources."
        ws["A2"].font = SUB
        ws.column_dimensions["A"].width = 60
        ws.column_dimensions["B"].width = 16
        ws.column_dimensions["C"].width = 90
        for i, h in enumerate(["Metric", "Value", "Definition / benchmark"], 1):
            c = ws.cell(row=4, column=i, value=h)
            c.font, c.fill = HDR, HDR_FILL
        yr = rng("Revenue", 5)
        P = "P&L & Cash"
        sumy = lambda sheet, row: f"SUMIF({yr},in_UEYear,{rng(sheet, row)})"
        items = [
            ("UE_Rev", "Revenue in the year", f"={sumy('Revenue', 28)}", EUR, ""),
            ("UE_CustMonths", "Customer-months in the year", f"={sumy('Revenue', 21)}", NUM, "sum of end-of-month paying customers"),
            ("ARPA", "ARPA (average revenue per account)", "=IF(in_UE_CustMonths>0,in_UE_Rev/in_UE_CustMonths,0)", EUR2,
             "EUR per paying company per month (flat price per company, not per employee)"),
            ("VarCost", "Variable cost per customer", "=in_InfraVar+in_PayFee*in_ARPA", EUR2, "variable infrastructure + payment fees"),
            ("CM", "Contribution margin (per-customer gross margin)", "=IF(in_ARPA>0,(in_ARPA-in_VarCost)/in_ARPA,0)", PCT,
             "Gross margin used in LTV. Excludes fixed hosting and support salary."),
            ("GMPL", "Gross margin, P&L basis", f"=IF(in_UE_Rev>0,{sumy(P, 10)}/in_UE_Rev,0)", PCT,
             "After fixed hosting and the onboarding/support salary - low while the base is small."),
            ("UE_SnM", "Sales & marketing spend in the year", f"={sumy(P, 13)}", EUR, "partner manager + marketing + travel + accountant commissions"),
            ("UE_New", "New paying customers in the year", f"={sumy('Revenue', 18)}", NUM, ""),
            ("CAC", "Blended CAC", "=IF(in_UE_New>0,in_UE_SnM/in_UE_New,0)", EUR, "S&M spend / new paying customers (fully loaded, incl. salary)"),
            ("CACPaid", "Paid-ads CAC (media only)", "=in_CostPerTrial/in_TrialConv", EUR, "cost per trial / trial conversion"),
            ("ChurnUE", "Monthly logo churn", "=in_Churn", PCT, "Benchmark: ChartMogul median 6.1%/month for ARPA < US$25"),
            ("Lifetime", "Expected customer lifetime", "=IF(in_Churn>0,1/in_Churn,0)", MON, "1 / churn"),
            ("LTV", "LTV (contribution basis)", "=IF(in_Churn>0,in_ARPA*in_CM/in_Churn,0)", EUR, "ARPA x contribution margin / churn"),
            ("LTVCAC", "LTV / CAC", "=IF(in_CAC>0,in_LTV/in_CAC,0)", MULT, "> 3x usually considered healthy"),
            ("Payback", "CAC payback", "=IF(in_ARPA*in_CM>0,in_CAC/(in_ARPA*in_CM),0)", '0.0" mo"',
             "Benchmark SMB SaaS: median ~11 months, healthy 6-14; 3-9 typical below US$5k ACV, so ours is above that (Aleph, secondary source)"),
        ]
        for i, (key, label, f, fmt, note) in enumerate(items):
            rr = 5 + i
            ws[f"A{rr}"], ws[f"B{rr}"], ws[f"C{rr}"] = label, f, note
            ws[f"B{rr}"].number_format = fmt
            ws[f"C{rr}"].font = SUB
            self.name(key, "Unit economics", f"B{rr}")

    # -- Use of funds ------------------------------------------------------------------
    def build_use_of_funds(self):
        ws = self.wb["Use of funds"]
        ws["A1"] = "Use of funds - pre-seed round"
        ws["A1"].font = TITLE
        ws["A2"] = ("Split = share of each category in planned spend over the zero-revenue runway (the months the round alone "
                    "pays for), applied to the round. Revenue earned in that period extends the runway further.")
        ws["A2"].font = SUB
        ws.column_dimensions["A"].width = 58
        ws.column_dimensions["B"].width = 16
        ws.column_dimensions["C"].width = 16
        ws.column_dimensions["D"].width = 12
        ws.column_dimensions["E"].width = 80
        ws["A4"], ws["B4"] = "Round size", "=in_Round"
        ws["B4"].number_format = EUR
        ws["B4"].font = GREEN
        ws["A5"], ws["B5"] = "Spend horizon = zero-revenue runway", "=in_RunwayZeroRev"
        ws["B5"].number_format = MON
        ws["B5"].font = GREEN
        ws["A6"], ws["B6"] = "  ending in", f"=INDEX({rng('Revenue', 4)},1,B5)"
        ws["B6"].number_format = "mmm yyyy"
        for i, h in enumerate(["Category", "Planned spend in horizon", "Allocated from round", "%", "What it buys"], 1):
            c = ws.cell(row=8, column=i, value=h)
            c.font, c.fill = HDR, HDR_FILL
        mo = rng("Revenue", 3)
        within = lambda row: f"SUMPRODUCT(({mo}<=$B$5)*{rng('Costs', row)})"
        R = lambda k: self.r("Costs", k)
        cats = [
            ("Product & engineering (2 co-founders full time)", [R("founder1"), R("founder2"), R("dev")],
             "Both co-founders full time on lean salaries; the developer is hired after the seed (in_DevStart), outside this horizon in the base case. "
             "REGES-Online, payroll validation fixes, billing, mobile store release."),
            ("Accountant channel & sales (partner manager, travel, commissions)", [R("sales"), R("travel"), R("comm")],
             "Partner manager (half time in year 1, full time from in_SalesFullStart) signing accounting practices; 20% x 6-month commissions."),
            ("Customer onboarding & support", [R("support")],
             "Half-time specialist from in_SupportStart: employee import, company setup and accounting mapping for each new firm."),
            ("Marketing (content, SEO, performance ads)", [R("mkt")], "REGES-Online / timesheet keywords; accountant materials."),
            ("Legal, payroll validation, compliance & security", [R("legal"), R("payrollval"), R("a11y"), R("pentest"), R("insurance")],
             "Round documents, DPA, payroll legal values validated, EN 301 549 accessibility audit, pentest, liability insurance."),
            ("Infrastructure & tools", [R("infrafix"), R("infravar"), R("tools"), R("payfee")], "Hosting, separate production database, AI/dev tools, payment fees."),
            ("Operations & contingency", [R("office"), R("admin"), R("contingency")], "Coworking in Timiș, bookkeeping, 5% contingency."),
        ]
        first = 9
        for i, (label, rows, what) in enumerate(cats):
            rr = first + i
            ws[f"A{rr}"] = label
            ws[f"B{rr}"] = "=" + "+".join(within(r_) for r_ in rows)
            ws[f"B{rr}"].number_format = EUR
            ws[f"C{rr}"] = f"=IF($B${first + len(cats)}>0,B{rr}/$B${first + len(cats)}*$B$4,0)"
            ws[f"C{rr}"].number_format = EUR
            ws[f"D{rr}"] = f"=IF($B$4>0,C{rr}/$B$4,0)"
            ws[f"D{rr}"].number_format = PCT
            ws[f"E{rr}"] = what
            ws[f"E{rr}"].font = SUB
        tr = first + len(cats)
        self.uof_rows = (first, tr - 1)
        ws[f"A{tr}"] = "Total"
        ws[f"A{tr}"].font = BOLD
        for cc, f, fmt in (("B", f"=SUM(B{first}:B{tr - 1})", EUR), ("C", f"=SUM(C{first}:C{tr - 1})", EUR), ("D", f"=SUM(D{first}:D{tr - 1})", PCT)):
            ws[f"{cc}{tr}"] = f
            ws[f"{cc}{tr}"].number_format = fmt
            ws[f"{cc}{tr}"].font = BOLD
            ws[f"{cc}{tr}"].border = TOP
        j = tr + 2
        self.section(ws, j, "Round size test (selected scenario; same pre-money valuation)", 5)
        for i, h in enumerate(["Round size tested", "EUR", "Runway, zero revenue", "Runway on plan", "Dilution at this pre-money"], 1):
            c = ws.cell(row=j + 1, column=i, value=h)
            c.font, c.fill = HDR, HDR_FILL
        cumcost, runmin = rng("P&L & Cash", 24), rng("P&L & Cash", 30)
        tests = [("VV ticket alone (not allowed without private match)", "=in_VVTicket"),
                 ("Proposed round", "=in_Round"),
                 ("VV maximum ticket + minimum 10% private match", "=200000*(1+in_MinPrivatePct)"),
                 ("Larger round (more angels)", 250000),
                 ("Upper end of founders' range", 300000)]
        for i, (lbl, v) in enumerate(tests):
            rr = j + 2 + i
            ws[f"A{rr}"], ws[f"B{rr}"] = lbl, v
            ws[f"B{rr}"].number_format = EUR
            ws[f"B{rr}"].font = GREEN if isinstance(v, str) else BLUE
            ws[f"C{rr}"] = f'=COUNTIF({cumcost},"<="&(B{rr}+in_OpeningCash))'
            ws[f"D{rr}"] = (f'=IF(COUNTIF({runmin},">="&-(B{rr}+in_OpeningCash))>={N_MONTHS},"> {N_MONTHS}",'
                            f'COUNTIF({runmin},">="&-(B{rr}+in_OpeningCash)))')
            ws[f"E{rr}"] = f"=B{rr}/(in_PreMoney+B{rr})"
            ws[f"C{rr}"].number_format = MON
            ws[f"D{rr}"].number_format = MON
            ws[f"E{rr}"].number_format = PCT
            if lbl == "Proposed round":
                ws[f"A{rr}"].font = BOLD
        self.round_test_rows = (j + 2, j + 6)
        ws[f"A{j + 7}"] = ("All rows use the same (lean) cost plan and the selected scenario. Runway on plan assumes the money arrives in model "
                           "month 1 (running minimum of cumulative operating cash flow). Dilution shown for the same pre-money; a different "
                           "round size would be negotiated at a different valuation.")
        ws[f"A{j + 7}"].font = SUB
        t220 = j + 4
        j = j + 9
        self.section(ws, j, "Why this round size", 5)
        lines = [
            '="Round: €"&TEXT(in_Round,"#,##0")&" = Vest Ventures accelerator ticket €"&TEXT(in_VVTicket,"#,##0")&" + independent angels €"&TEXT(in_Angels,"#,##0")&"."',
            '="Zero-revenue runway: "&in_RunwayZeroRev&" months (option A, VV + minimum 10% match = €"&TEXT(in_OptionA,"#,##0")&": "&in_RunwayZeroRevA&" months)."',
            ('=IF(in_CashLow<0,"On the selected scenario cash lasts "&in_RunwayPlan&" months on plan without further funding; the seed round '
             'is raised from the 100-paying-customer milestone (model month "&in_M100&").","On the selected scenario cash never falls below €"'
             '&TEXT(in_CashLow,"#,##0")&"; break-even in model month "&in_Breakeven&".")'),
            '="Next milestone (seed readiness) inside the runway: 100 paying firms by model month "&in_M100&", payroll validated, REGES-Online in daily use, accountant channel repeatable."',
            ('="Why this size: the round funds the path to the seed trigger (100 paying firms, model month "&in_M100&") with "&in_RunwayPlan&'
             '" months of runway on plan and "&TEXT(in_InvPct,"0.0%")&" dilution, on a lean cost plan: co-founders at €"&TEXT(in_SalFounder,"#,##0")&'
             '" gross, partner manager at "&TEXT(in_SalesFTE,"0%")&" until model month "&in_SalesFullStart&" (partner signings scaled to match), '
             'support specialist at "&TEXT(in_SupportFTE,"0%")&" from model month "&in_SupportStart&", developer after the seed (model month "&in_DevStart&")."'),
            (f'="Larger rounds (table above) buy runway, e.g. €"&TEXT(B{t220},"#,##0")&" (VV maximum ticket + minimum match): "&C{t220}&'
             f'" months with zero revenue, "&D{t220}&" on plan - but need a larger VV ticket and more private money, and no angel is identified yet."'),
        ]
        for k, f in enumerate(lines):
            ws[f"A{j + 1 + k}"] = f

    # -- Cap table ---------------------------------------------------------------------
    def build_cap_table(self):
        ws = self.wb["Cap table"]
        ws["A1"] = "Cap table - today and after the pre-seed round (fully diluted)"
        ws["A1"].font = TITLE
        ws["A2"] = ("Today: SRL, 2 co-founders, no investors, no options (founders' statement). Units below are NOTIONAL (10,000 = 100%); "
                    "the SRL's real number of social parts is [TO CONFIRM]. Instrument: Vest Ventures templates are a CLA or an SHA "
                    "(vestventures.vc/en/programs-terms); modelled as priced equity at the pre-money valuation.")
        ws["A2"].font = SUB
        ws.column_dimensions["A"].width = 52
        for cc in "BCDEFG":
            ws.column_dimensions[cc].width = 16
        ws["A4"], ws["B4"] = "Pre-money valuation (fully diluted)", "=in_PreMoney"
        ws["A5"], ws["B5"] = "Round size", "=in_Round"
        ws["A6"], ws["B6"] = "Post-money valuation", "=B4+B5"
        ws["A7"], ws["B7"] = "Investors' ownership post-money", "=B5/B6"
        ws["A8"], ws["B8"] = "ESOP pool, % post-money (created pre-money)", "=in_ESOP"
        ws["A9"], ws["B9"] = "Founders' units today (notional)", 10000
        ws["A10"], ws["B10"] = "Fully diluted units post-money", "=B9/(1-B7-B8)"
        ws["A11"], ws["B11"] = "Price per unit (pre-money / pre-money FD units)", "=B4/(B9+B8*B10)"
        for rr, fmt in zip(range(4, 12), (EUR, EUR, EUR, PCT, PCT, NUM, NUM, EUR2)):
            ws[f"B{rr}"].number_format = fmt
            ws[f"B{rr}"].font = GREEN if isinstance(ws[f"B{rr}"].value, str) and "in_" in ws[f"B{rr}"].value else BLACK
        ws["B9"].font = BLUE
        self.name("PostMoney", "Cap table", "B6")
        self.name("InvPct", "Cap table", "B7")
        heads = ["Holder", "Today - units", "Today - %", "Pre-seed - units", "Pre-seed - %", "Invested (EUR)"]
        for i, h in enumerate(heads, 1):
            c = ws.cell(row=13, column=i, value=h)
            c.font, c.fill = HDR, HDR_FILL
        holders = [
            ("Miroslav Maletici (co-founder)", "=0.51*$B$9", None),
            ("Răzvan Pervulescu (co-founder)", "=0.49*$B$9", None),
            ("ESOP pool (unallocated)", 0, "=$B$8*$B$10"),
            ("Vest Ventures (accelerator ticket)", 0, "=in_VVTicket/$B$11"),
            ("Independent angels", 0, "=in_Angels/$B$11"),
        ]
        for i, (h, today, post) in enumerate(holders):
            rr = 14 + i
            ws[f"A{rr}"] = h
            ws[f"B{rr}"] = today
            if i < 2:
                ws[f"B{rr}"].font = BLUE
                ws[f"D{rr}"] = f"=B{rr}"
            else:
                ws[f"D{rr}"] = post
            ws[f"C{rr}"] = f"=B{rr}/$B$19"
            ws[f"E{rr}"] = f"=D{rr}/$D$19"
            ws[f"F{rr}"] = {3: "=in_VVTicket", 4: "=in_Angels"}.get(i, 0)
            for cc, fmt in (("B", NUM), ("C", PCT), ("D", NUM), ("E", PCT), ("F", EUR)):
                ws[f"{cc}{rr}"].number_format = fmt
        ws["A19"] = "Total"
        for cc, f, fmt in (("B", "=SUM(B14:B18)", NUM), ("C", "=SUM(C14:C18)", PCT), ("D", "=SUM(D14:D18)", NUM),
                           ("E", "=SUM(E14:E18)", PCT), ("F", "=SUM(F14:F18)", EUR)):
            ws[f"{cc}19"] = f
            ws[f"{cc}19"].number_format = fmt
            ws[f"{cc}19"].font = BOLD
            ws[f"{cc}19"].border = TOP
        ws["A19"].font = BOLD
        ws["A20"] = "Check: post-money of the round at unit price (must equal post-money)"
        ws["D20"] = "=D19*B11"
        ws["D20"].number_format = EUR
        self.name("FoundersPostM", "Cap table", "E14")
        self.name("FoundersPostR", "Cap table", "E15")
        self.name("VVPct", "Cap table", "E17")
        self.name("AngelPct", "Cap table", "E18")
        self.name("ESOPPct", "Cap table", "E16")
        notes = [
            "Convertible alternative (CLA / SAFE-like): Vest Ventures' default instruments are a convertible loan agreement or an SHA.",
            "A CLA would carry a valuation cap equal to the pre-money above and a discount (e.g. 20%) to the next priced round; the CLA "
            "conversion math is on cap-table.xlsx, sheet 'CLA alternative'.",
            "SRL caveat: changing an SRL's articles needs every associate's vote (Companies Law 31/1990, art. 192); conversion SRL -> SA or a holding "
            "at the round is [TO CONFIRM] with a lawyer. An ESOP in an SRL is usually a virtual / phantom scheme until then.",
        ]
        for i, t in enumerate(notes):
            ws[f"A{22 + i}"] = t
            ws[f"A{22 + i}"].font = SUB

    # -- Sources -----------------------------------------------------------------------
    def build_sources(self):
        ws = self.wb["Sources"]
        ws["A1"] = "Sources"
        ws["A1"].font = TITLE
        ws.column_dimensions["A"].width = 60
        ws.column_dimensions["B"].width = 120
        rows = [
            ("Pricing, modules, accountant pilot, zero customers", "https://administrativo.ro/en/preturi ; https://administrativo.ro/pentru-contabili ; product code and founders' statement (zero paying customers on 6 Oct 2026)"),
            ("Monthly churn benchmark 6.1% (ARPA < US$25)", "ChartMogul - https://chartmogul.com/blog/good-customer-churn-rate/"),
            ("SMB CAC payback ~11 months median, 6-14 healthy; 3-9 typical below US$5k ACV","Aleph - https://www.getaleph.com/answers/cac-payback-period-saas-2026 (secondary)"),
            ("Romania 2025: EUR 103M VC in 40 deals; avg pre-seed EUR 526k; avg seed EUR 1.5M",
             "How to Web & Underline Ventures via Romania Insider - https://www.romania-insider.com/romanian-startups-funding-2025"),
            ("Vest Ventures stages, tickets (Accelerator EUR 10k-200k, >=10% private co-investment), instruments (CLA/SHA)",
             "https://vestventures.vc/en/stages-accelerator ; https://vestventures.vc/en/programs-terms ; https://vestventures.vc/en/faq"),
            ("Employers with >=1 employee: ~526,000 (FY2025)", "ICAP CRIF via Curierul Național - https://curierulnational.ro/peste-jumatate-de-milion-de-firme-aveau-388-milioane-salariati-in-2025/"),
            ("Enterprises by size class (Eurostat SBS 2023)", "https://ec.europa.eu/eurostat/databrowser/view/sbs_sc_ovw/default/table"),
            ("Accounting practices ~18,400 (DERIVED: 12,776 firms + 5,671 individuals, CECCAR 31 Dec 2024)",
             "https://ceccar.ro/ro/wp-content/uploads/2025/08/PDF-Raport-CECCAR-2025-evaluare-risc-spalarea-banilor.pdf (search excerpt; PDF to open by hand)"),
            ("Internal cost estimates (legal, insurance, infra, tools, payroll validation EUR 8k, pentest EUR 6k)", "Founders' internal cost estimates (not externally sourced)"),
            ("Payroll legal values to be confirmed (incl. CAM)", "Product documentation: every payroll legal value is marked for accountant validation"),
            ("Salaries, conversion rates, partner productivity, plan mix, marketing budgets", "ASSUMPTIONS by the founders' model - no external source; see Assumptions!D"),
        ]
        for i, h in enumerate(["What", "Source"], 1):
            c = ws.cell(row=3, column=i, value=h)
            c.font, c.fill = HDR, HDR_FILL
        for i, (a, b) in enumerate(rows):
            ws[f"A{4 + i}"], ws[f"B{4 + i}"] = a, b

    # -- Cover -------------------------------------------------------------------------
    def build_cover(self):
        ws = self.cover
        ws.column_dimensions["A"].width = 4
        ws.column_dimensions["B"].width = 58
        ws.column_dimensions["C"].width = 22
        ws.column_dimensions["D"].width = 70
        ws["B2"] = "ADMINISTRATIVO - financial model (pre-seed)"
        ws["B2"].font = TITLE
        ws["B3"] = "HR & administration SaaS for Romanian SMEs - administrativo.ro"
        ws["B4"], ws["C4"] = "Prepared", "7 Oct 2026"
        ws["B5"], ws["C5"] = "Currency", "EUR (RON converted at 5.0 RON = 1 EUR, Assumptions!in_FX)"
        ws["B6"], ws["C6"] = "Status on 6 Oct 2026", "Product live since early Sept 2026; ZERO paying customers; free pilots only"
        ws["B8"] = "SCENARIO SELECTOR (1 = Base, 2 = Conservative, 3 = Upside)"
        ws["B8"].font = BOLD
        ws["C8"] = self.scenario
        ws["C8"].font = BLUE_B
        ws["C8"].fill = YELLOW
        ws["D8"] = "=INDEX(Scenarios!$C$4:$E$4,1,C8)"
        ws["D8"].font = GREEN
        self.wb.defined_names["ScenarioIdx"] = DefinedName("ScenarioIdx", attr_text="Cover!$C$8")
        ws["B10"] = "Key outputs (selected scenario, PROJECTION)"
        ws["B10"].font = BOLD
        P = "P&L & Cash"
        outs = [
            ("Round size", "=in_Round", EUR),
            ("Pre-money / post-money valuation", '="€"&TEXT(in_PreMoney,"#,##0")&" / €"&TEXT(in_PostMoney,"#,##0")', None),
            ("Investors' ownership post-money", "=in_InvPct", PCT),
            ("Model month 1 (assumed close)", "=in_StartDate", "mmm yyyy"),
            ("Runway with zero revenue", "=in_RunwayZeroRev", MON),
            ("Runway on plan", "=in_RunwayPlan", MON),
            ("Paying customers, month 12 / 24 / 36",
             f'=TEXT(Revenue!{col(12)}21,"0")&" / "&TEXT(Revenue!{col(24)}21,"0")&" / "&TEXT(Revenue!{col(36)}21,"0")', None),
            ("MRR month 12", f"=Revenue!{col(12)}26", EUR),
            ("MRR month 24", f"=Revenue!{col(24)}26", EUR),
            ("MRR month 36", f"=Revenue!{col(36)}26", EUR),
            ("ARR month 36", f"=Revenue!{col(36)}27", EUR),
            ("Break-even month (model month #)", "=in_Breakeven", NUM),
            ("Cash low point", "=in_CashLow", EUR),
            ("Months to 30 / 100 / 300 paying customers", '=in_M30&" / "&in_M100&" / "&in_M300', None),
            ("Blended ARPA (year 2) / LTV / CAC", '="€"&TEXT(in_ARPA,"0")&" / €"&TEXT(in_LTV,"#,##0")&" / €"&TEXT(in_CAC,"#,##0")', None),
            ("LTV/CAC, CAC payback", '=TEXT(in_LTVCAC,"0.0")&"x, "&TEXT(in_Payback,"0.0")&" months"', None),
            ("Integrity checks (plan mix = 100%, EBITDA check = 0)", '=IF(AND(ROUND(in_MixCheck,6)=1,in_MaxCheck=0),"OK","CHECK")', None),
        ]
        for i, (label, f, fmt) in enumerate(outs):
            rr = 11 + i
            ws[f"B{rr}"], ws[f"C{rr}"] = label, f
            ws[f"C{rr}"].font = GREEN
            if fmt:
                ws[f"C{rr}"].number_format = fmt
        rr = 11 + len(outs) + 1
        ws[f"B{rr}"] = "Legend"
        ws[f"B{rr}"].font = BOLD
        leg = [("Blue text", "hard-coded input - edit on Assumptions / Scenarios", BLUE),
               ("Green text", "link to another sheet", GREEN),
               ("Black text", "formula", BLACK),
               ("Yellow fill", "key lever (scenario, round, valuation, ESOP)", BLACK)]
        for i, (a, b, f) in enumerate(leg):
            ws[f"B{rr + 1 + i}"], ws[f"C{rr + 1 + i}"] = a, b
            ws[f"B{rr + 1 + i}"].font = f
        ws[f"B{rr + 4}"].fill = YELLOW
        rr += 6
        ws[f"B{rr}"] = "Honesty notes"
        ws[f"B{rr}"].font = BOLD
        notes = [
            "All figures after 6 Oct 2026 are PROJECTIONS built bottom-up from pilots -> accountant partners -> direct trials. No revenue exists today.",
            "Payroll is built, undergoing accountant validation; the Finance package share is kept low.",
            "Prices above 20 employees are not published; larger firms are modelled at list price.",
            "No billing integration exists in code yet; payment fees are an assumption.",
            "Salaries, conversion rates and partner productivity are founders' assumptions (no external source).",
            "Lean pre-seed cost plan: co-founders on lean salaries, partner manager and support specialist part-time first "
            "(partner signings scaled to the manager's hours), developer hired only after the seed.",
            "[TO CONFIRM]: legal entity and registered office (VV requires Arad, Caraș-Severin, Hunedoara or Timiș), tax regime, social parts.",
            "Sheets: Assumptions -> Revenue -> Costs -> P&L & Cash -> Unit economics; Scenarios; Use of funds; Cap table; Sources.",
        ]
        for i, t in enumerate(notes):
            ws[f"B{rr + 1 + i}"] = t
            ws[f"B{rr + 1 + i}"].font = SUB


# --------------------------------------------------------------------------------------
# Standalone cap table workbook
# --------------------------------------------------------------------------------------

CAP = dict(founders_units=10000, m_pct=0.51, r_pct=0.49, esop=0.10, vv=135000, angels=15000, pre=1250000,
           seed_amount=1000000, seed_pre=4000000, seed_esop_topup=0.0, cla_discount=0.20, cla_interest=0.0, cla_years=2.0)


def build_cap_table_wb() -> Workbook:
    wb = Workbook()
    ws = wb.active
    ws.title = "Cap table"
    ws["A1"] = "ADMINISTRATIVO - pro-forma cap table (fully diluted, notional units)"
    ws["A1"].font = TITLE
    ws["A2"] = ("Today = founders' statement (SRL, 51/49, no investors, no options). Pre-seed = this round. Seed = ILLUSTRATIVE only "
                "(not a plan, not a commitment). Units are notional (10,000 = 100% today); the SRL's real social parts are [TO CONFIRM].")
    ws["A2"].font = SUB
    ws.column_dimensions["A"].width = 52
    for cc in "BCDEFGHI":
        ws.column_dimensions[cc].width = 15
    ws.column_dimensions["J"].width = 70
    inputs = [
        ("Founders' units today (notional)", CAP["founders_units"], NUM, "notional"),
        ("Miroslav Maletici share today", CAP["m_pct"], PCT, "founders' statement"),
        ("Răzvan Pervulescu share today", CAP["r_pct"], PCT, "founders' statement"),
        ("Pre-seed: Vest Ventures accelerator ticket (EUR)", CAP["vv"], EUR, "Vest Ventures Accelerator: EUR 10k-200k (vestventures.vc)"),
        ("Pre-seed: independent angels (EUR)", CAP["angels"], EUR, ">= 10% private co-investment rule"),
        ("Pre-seed: pre-money valuation, fully diluted (EUR)", CAP["pre"], EUR, "same as financial-model.xlsx"),
        ("Pre-seed: ESOP pool, % post-money, created pre-money", CAP["esop"], PCT, "dilutes founders only"),
        ("Seed (ILLUSTRATIVE): amount (EUR)", CAP["seed_amount"], EUR, "VV Seed range EUR 200k-1M; RO avg seed 2025 EUR 1.5M (How to Web & Underline Ventures)"),
        ("Seed (ILLUSTRATIVE): pre-money valuation (EUR)", CAP["seed_pre"], EUR, "ILLUSTRATIVE assumption, needs traction"),
        ("Seed (ILLUSTRATIVE): ESOP top-up, % post-seed", CAP["seed_esop_topup"], PCT, "0 = no top-up"),
    ]
    for i, h in enumerate(["Input", "Value", "", "", "", "", "", "", "", "Note"], 1):
        if h:
            c = ws.cell(row=4, column=i, value=h)
            c.font, c.fill = HDR, HDR_FILL
    for i, (lbl, v, fmt, note) in enumerate(inputs):
        rr = 5 + i
        ws[f"A{rr}"], ws[f"B{rr}"], ws[f"J{rr}"] = lbl, v, note
        ws[f"B{rr}"].font = BLUE
        ws[f"B{rr}"].number_format = fmt
        ws[f"J{rr}"].font = SUB
    for rr in (8, 9, 10, 11):
        ws[f"B{rr}"].fill = YELLOW
    # B5 units, B6 M, B7 R, B8 VV, B9 angels, B10 pre, B11 esop, B12 seed amt, B13 seed pre, B14 topup
    ws["A16"] = "Derived"
    ws["A16"].font = BOLD
    der = [
        ("Pre-seed round (EUR)", "=B8+B9", EUR),                                   # B17
        ("Pre-seed post-money (EUR)", "=B10+B17", EUR),                            # B18
        ("Pre-seed investors' % post", "=B17/B18", PCT),                           # B19
        ("Pre-seed FD units post", "=B5/(1-B19-B11)", NUM),                        # B20
        ("Pre-seed price per unit (EUR)", "=B10/(B5+B11*B20)", EUR2),              # B21
        ("Seed post-money (EUR)", "=B13+B12", EUR),                                # B22
        ("Seed investors' % post", "=B12/B22", PCT),                               # B23
        ("Seed FD units post", "=B20/(1-B23-B14)", NUM),                           # B24
        ("Seed price per unit (EUR)", "=B13/(B20+B14*B24)", EUR2),                 # B25
    ]
    for i, (lbl, f, fmt) in enumerate(der):
        rr = 17 + i
        ws[f"A{rr}"], ws[f"B{rr}"] = lbl, f
        ws[f"B{rr}"].number_format = fmt
    hdr = ["Holder", "Today units", "Today %", "Pre-seed units", "Pre-seed %", "Seed units", "Seed %", "Invested (EUR)"]
    for i, h in enumerate(hdr, 1):
        c = ws.cell(row=28, column=i, value=h)
        c.font, c.fill = HDR, HDR_FILL
    holders = [
        ("Miroslav Maletici (co-founder)", "=B5*B6", "=B29", "=D29"),
        ("Răzvan Pervulescu (co-founder)", "=B5*B7", "=B30", "=D30"),
        ("ESOP pool", "0", "=B11*B20", "=D31+B14*B24"),
        ("Vest Ventures", "0", "=B8/B21", "=D32"),
        ("Independent angels", "0", "=B9/B21", "=D33"),
        ("Seed investors (ILLUSTRATIVE)", "0", "0", "=B12/B25"),
    ]
    for i, (h, t, p, s) in enumerate(holders):
        rr = 29 + i
        ws[f"A{rr}"] = h
        ws[f"B{rr}"] = t if t.startswith("=") else 0
        ws[f"D{rr}"] = p if p.startswith("=") else 0
        ws[f"F{rr}"] = s
        ws[f"C{rr}"] = f"=B{rr}/$B$35"
        ws[f"E{rr}"] = f"=D{rr}/$D$35"
        ws[f"G{rr}"] = f"=F{rr}/$F$35"
        ws[f"H{rr}"] = {3: "=B8", 4: "=B9", 5: "=B12"}.get(i, 0)
        for cc, fmt in (("B", NUM), ("C", PCT), ("D", NUM), ("E", PCT), ("F", NUM), ("G", PCT), ("H", EUR)):
            ws[f"{cc}{rr}"].number_format = fmt
    ws["A35"] = "Total"
    ws["A35"].font = BOLD
    for cc, fmt in (("B", NUM), ("C", PCT), ("D", NUM), ("E", PCT), ("F", NUM), ("G", PCT), ("H", EUR)):
        ws[f"{cc}35"] = f"=SUM({cc}29:{cc}34)"
        ws[f"{cc}35"].number_format = fmt
        ws[f"{cc}35"].font = BOLD
        ws[f"{cc}35"].border = TOP
    ws["A36"] = "Founders combined"
    for cc in "CEG":
        ws[f"{cc}36"] = f"={cc}29+{cc}30"
        ws[f"{cc}36"].number_format = PCT
    ws["A37"] = "Check: units x price = post-money (pre-seed, seed)"
    ws["D37"], ws["F37"] = "=D35*B21", "=F35*B25"
    ws["D37"].number_format = ws["F37"].number_format = EUR

    # CLA alternative
    w2 = wb.create_sheet("CLA alternative")
    w2["A1"] = "Alternative: convertible loan agreement (CLA) instead of priced equity"
    w2["A1"].font = TITLE
    w2["A2"] = ("Vest Ventures contracts: term sheet, then a CLA or an SHA on the fund's templates (vestventures.vc/en/programs-terms). "
                "Interest, discount and cap below are ILLUSTRATIVE inputs, not VV terms [TO CONFIRM].")
    w2["A2"].font = SUB
    w2.column_dimensions["A"].width = 64
    w2.column_dimensions["B"].width = 18
    w2.column_dimensions["C"].width = 70
    rows = [
        ("CLA principal (EUR)", "='Cap table'!B17", EUR, "same money as the pre-seed round"),              # B4
        ("Valuation cap, pre-money FD (EUR)", "='Cap table'!B10", EUR, "cap = pre-money of the priced alternative"),  # B5
        ("Discount to the next priced round", CAP["cla_discount"], PCT, "ILLUSTRATIVE"),                     # B6
        ("Annual interest (simple)", CAP["cla_interest"], PCT, "ILLUSTRATIVE; VV template terms [TO CONFIRM]"),  # B7
        ("Years until conversion", CAP["cla_years"], '0.0', "ILLUSTRATIVE (seed ~24 months after close: 100-customer trigger + raise)"),  # B8
        ("Amount converting (EUR)", "=B4*(1+B7*B8)", EUR, ""),                                               # B9
        ("Pre-conversion units (founders + ESOP pool, notional)", "='Cap table'!B5+'Cap table'!D31", NUM,
         "same pool units as the priced round (10% of post-money), so the two paths are comparable"),  # B10
        ("Cap price per unit (EUR)", "=B5/B10", EUR2, "cap / pre-conversion units"),                         # B11
        ("Seed price per unit if no CLA (seed pre-money / pre-conversion units)", "='Cap table'!B13/B10", EUR2, "ILLUSTRATIVE seed"),  # B12
        ("Discounted seed price", "=B12*(1-B6)", EUR2, ""),                                                  # B13
        ("Conversion price = MIN(cap price, discounted price)", "=MIN(B11,B13)", EUR2, ""),                  # B14
        ("CLA units at conversion", "=B9/B14", NUM, ""),                                                     # B15
        ("CLA holders' % right after conversion (before seed money)", "=B15/(B10+B15)", PCT, ""),            # B16
        ("Seed amount (EUR)", "='Cap table'!B12", EUR, ""),                                                  # B17
        ("Seed new units (at seed price, after conversion)", "=B17/('Cap table'!B13/(B10+B15))", NUM,
         "pre-money includes converted units (simplification)"),                                             # B18
        ("Total units post-seed", "=B10+B15+B18", NUM, ""),                                                  # B19
        ("Founders combined after seed - CLA path", "=('Cap table'!B5)/B19", PCT, ""),                       # B20
        ("Founders combined after seed - priced pre-seed path", "='Cap table'!G36", PCT, "from sheet 'Cap table'"),  # B21
        ("CLA holders after seed", "=B15/B19", PCT, ""),                                                     # B22
    ]
    for i, (lbl, v, fmt, note) in enumerate(rows):
        rr = 4 + i
        w2[f"A{rr}"], w2[f"B{rr}"], w2[f"C{rr}"] = lbl, v, note
        w2[f"B{rr}"].number_format = fmt
        w2[f"C{rr}"].font = SUB
        if not (isinstance(v, str) and v.startswith("=")):
            w2[f"B{rr}"].font = BLUE
        elif "Cap table" in v:
            w2[f"B{rr}"].font = GREEN
    w3 = wb.create_sheet("Notes")
    notes = [
        "Ownership today: Miroslav Maletici 51%, Răzvan Pervulescu 49% - founders' statement; no investors, no options issued.",
        "Legal entity named on administrativo.ro as its operator: WISELEARNING S.R.L. (CUI 50321210, J35/2618/2024) - [TO CONFIRM] that this SRL holds the 51/49 cap table.",
        "Pre-seed modelled as priced equity at EUR 1.25M pre-money, fully diluted including a new 10% ESOP pool created pre-money.",
        "Vest Ventures may use a CLA (convertible loan) or an SHA; see 'CLA alternative'.",
        "SRL governance: amendments to the articles need unanimity of associates; an SRL -> SA or holding conversion at the round is [TO CONFIRM] with a lawyer.",
        "Seed column is ILLUSTRATIVE only: no seed investor, amount or valuation has been discussed.",
        "Same figures as financial-model.xlsx, sheet 'Cap table'.",
    ]
    w3.column_dimensions["A"].width = 150
    for i, t in enumerate(notes):
        w3[f"A{1 + i}"] = t
    for w in wb.worksheets:
        style_all(w)
        w.sheet_view.showGridLines = False
    wb.calculation.fullCalcOnLoad = True
    return wb


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--scenario", type=int, default=1)
    ap.add_argument("--out", default=str(ROOT / "financial-model.xlsx"))
    ap.add_argument("--snapshot", default=None, help="JSON with per-scenario results to paste on 'Scenarios'")
    ap.add_argument("--cap-table", default=str(ROOT / "cap-table.xlsx"))
    ap.add_argument("--no-cap-table", action="store_true")
    a = ap.parse_args()
    snap = json.loads(Path(a.snapshot).read_text()) if a.snapshot else None
    Model(a.scenario, snap).build().save(a.out)
    print("wrote", a.out)
    if not a.no_cap_table:
        build_cap_table_wb().save(a.cap_table)
        print("wrote", a.cap_table)


if __name__ == "__main__":
    main()
